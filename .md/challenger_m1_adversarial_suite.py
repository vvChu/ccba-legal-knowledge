"""
Comprehensive Adversarial Verification Suite for Milestone 1
Base QCVN 06:2022/BXD Parity & Integrity Verification
Author: challenger_m1_1
"""
import sys
import os
import re
import json
import difflib
from pathlib import Path
import docx

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

DOCX_PATH = Path(r"d:\GitHubProjects\ccba-legal-knowledge\.md\extracted_docs\qcvn_06_2022_bxd\qcvn_06_2022_bxd.docx")
MD_PATH = Path(r"d:\GitHubProjects\ccba-legal-knowledge\legal_docs\02_qcvn\qcvn_06_2022_bxd\qcvn_06_2022_bxd.md")
OUT_JSON = Path(r"d:\GitHubProjects\ccba-legal-knowledge\.md\challenger_m1_adversarial_results.json")
OUT_REPORT = Path(r"d:\GitHubProjects\ccba-legal-knowledge\.md\challenger_m1_adversarial_report.md")

def normalize_text(text: str) -> str:
    if not text:
        return ""
    # Normalize unicode whitespace, non-breaking spaces, zero-width characters
    text = text.replace('\xa0', ' ').replace('\u200b', '').replace('\ufeff', '')
    # Normalize markdown bold, italics, code markers, headers, anchors
    text = re.sub(r'<a id="[^"]+"></a>', '', text)
    text = re.sub(r'[#\*_`]', '', text)
    # Normalize multiple whitespace
    text = re.sub(r'\s+', ' ', text).strip()
    return text

def strip_all_punct(text: str) -> str:
    return re.sub(r'[^\w\d]+', '', text.lower())

def run_suite():
    print("=" * 80)
    print("RUNNING CHALLENGER M1 ADVERSARIAL VERIFICATION SUITE")
    print("=" * 80)

    doc = docx.Document(str(DOCX_PATH))
    md_raw = MD_PATH.read_text(encoding="utf-8")
    md_lines = md_raw.splitlines()

    print(f"Loaded DOCX: {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables.")
    print(f"Loaded MD: {len(md_lines)} lines, {len(md_raw)} characters.")

    # 1. Classify all docx paragraphs
    # In QCVN 06:2022/BXD docx:
    # - Paragraphs 0..21 are cover page, title, TOC (Mục lục), etc.
    # - Paragraph 22 is 'LỜI NÓI ĐẦU' or '1. QUY ĐỊNH CHUNG'
    all_docx_paras = []
    for idx, p in enumerate(doc.paragraphs):
        raw = p.text.strip()
        norm = normalize_text(raw)
        punct_strip = strip_all_punct(raw)
        if norm:
            all_docx_paras.append({
                "docx_idx": idx,
                "raw": raw,
                "norm": norm,
                "punct_strip": punct_strip,
                "length": len(norm),
                "style": p.style.name if p.style else None
            })

    print(f"Total non-empty DOCX paragraphs: {len(all_docx_paras)}")

    # 2. Extract MD paragraphs / elements
    # Clean md text
    clean_md_global = normalize_text(md_raw)
    punct_md_global = strip_all_punct(md_raw)

    # 3. Test 1: Full-length normalized exact paragraph search
    # Check what matches 100% full string vs partial vs missing
    matched_full = []
    matched_punct = []
    matched_fuzzy = []
    unmatched = []

    for item in all_docx_paras:
        norm = item["norm"]
        punct = item["punct_strip"]

        if norm in clean_md_global:
            matched_full.append(item)
        elif punct and punct in punct_md_global:
            matched_punct.append(item)
        else:
            # Check fuzzy / longest common substring
            # Search for best match in MD lines
            best_ratio = 0
            best_candidate = ""
            for line in md_lines:
                norm_line = normalize_text(line)
                if not norm_line:
                    continue
                ratio = difflib.SequenceMatcher(None, norm, norm_line).ratio()
                if ratio > best_ratio:
                    best_ratio = ratio
                    best_candidate = norm_line

            if best_ratio >= 0.85:
                matched_fuzzy.append({**item, "best_ratio": best_ratio, "best_candidate": best_candidate})
            else:
                unmatched.append({**item, "best_ratio": best_ratio, "best_candidate": best_candidate})

    print("\n--- PARAGRAPH RETENTION AUDIT ---")
    print(f"1. Exact Normalized Matches: {len(matched_full)} / {len(all_docx_paras)} ({len(matched_full)/len(all_docx_paras)*100:.2f}%)")
    print(f"2. Punctuation-Stripped Exact Matches: {len(matched_punct)} / {len(all_docx_paras)} ({len(matched_punct)/len(all_docx_paras)*100:.2f}%)")
    print(f"3. High-Fuzzy (>=85% similarity) Matches: {len(matched_fuzzy)} / {len(all_docx_paras)} ({len(matched_fuzzy)/len(all_docx_paras)*100:.2f}%)")
    print(f"4. Unmatched / Questionable Items: {len(unmatched)}")

    # 4. Detailed Analysis of Unmatched and Fuzzy items
    print("\n--- DETAILED INSPECTION OF NON-EXACT MATCHES ---")
    for idx, item in enumerate(unmatched):
        print(f"\n[UNMATCHED #{idx+1}] Docx P#{item['docx_idx']} (Len {item['length']}):")
        print(f"  DOCX : {item['raw']}")
        print(f"  BEST MATCH (ratio {item.get('best_ratio', 0):.2f}): {item.get('best_candidate', '')[:120]}")

    for idx, item in enumerate(matched_fuzzy):
        print(f"\n[FUZZY #{idx+1}] Docx P#{item['docx_idx']} (Ratio: {item['best_ratio']:.2f}):")
        print(f"  DOCX: {item['raw']}")
        print(f"  MD  : {item['best_candidate'][:120]}")

    # 5. Heading & Numbering Anomaly Audits
    print("\n" + "=" * 80)
    print("HEADING & NUMBERING ANOMALY AUDIT")
    print("=" * 80)

    # A. Section 1.4 definition numbers (1.4.1 to 1.4.72)
    sec_1_4_headings = re.findall(r'^(#{1,6}\s+1\.4\.(\d+)(?:\.|\s+|$).*)', md_raw, re.MULTILINE)
    sec_1_4_dot_split = re.findall(r'^(#{1,6}\s+1\.4\.\d+\.\d+.*)', md_raw, re.MULTILINE)
    print(f"Section 1.4 regular definitions found: {len(sec_1_4_headings)}")
    print(f"Section 1.4 dot-split bugs (e.g. 1.4.1.7) found: {len(sec_1_4_dot_split)}")
    if sec_1_4_dot_split:
        print(f"  FAIL: Found {len(sec_1_4_dot_split)} dot-split bugs: {sec_1_4_dot_split[:5]}")
    else:
        print("  PASS: 0 dot-split bugs in Section 1.4.")

    # Check 1.4 definitions completeness from 1 to 72
    def_numbers = [int(num) for _, num in sec_1_4_headings]
    missing_defs = [i for i in range(1, 73) if i not in def_numbers]
    print(f"Section 1.4 definition sequence (1..72): Missing numbers: {missing_defs}")

    # B. Space headings in Chapter 4 and 7
    space_h4 = re.findall(r'^#{1,6}\s+4\s+\d+.*', md_raw, re.MULTILINE)
    space_h7 = re.findall(r'^#{1,6}\s+7\s+\d+.*', md_raw, re.MULTILINE)
    print(f"Chapter 4 space headings (`### 4 X`): {len(space_h4)}")
    print(f"Chapter 7 space headings (`### 7 X`): {len(space_h7)}")

    # C. Missing subdots in Chapters 2, 5, 6
    subdot_targets = [
        "2.1.1.1", "2.1.1.2", "2.2.1.1", "2.2.1.2", "2.2.1.3", "2.5.6.3.3",
        "5.1.1.1", "5.1.1.2", "5.1.1.3", "5.1.1.4",
        "6.2.1.1", "6.2.1.2", "6.2.1.3", "6.2.1.4"
    ]
    missing_subdots = []
    for sd in subdot_targets:
        if not re.search(r'#{1,6}\s+' + re.escape(sd), md_raw):
            missing_subdots.append(sd)
    print(f"Missing subdot headings count: {len(missing_subdots)} (Missing: {missing_subdots})")

    # D. Residual table cell headings (e.g. ### 24 0, ### 1 4, ### 0 5, etc.)
    fake_table_headings = re.findall(r'^#{1,6}\s+\d+\s+\d+\b.*', md_raw, re.MULTILINE)
    print(f"Residual fake space headings (e.g. `### 24 0`): {len(fake_table_headings)}")
    if fake_table_headings:
        print(f"  Residual fake headings: {fake_table_headings[:10]}")

    # E. Heading level distribution
    h_dist = {}
    for line in md_lines:
        m = re.match(r'^(#{1,6})\s+', line)
        if m:
            lvl = len(m.group(1))
            h_dist[lvl] = h_dist.get(lvl, 0) + 1
    print(f"Heading level distribution: {sorted(h_dist.items())}")

    # F. Check Appendix headings (Phụ lục A to I)
    appendix_letters = ["A", "B", "C", "D", "E", "F", "G", "H", "I"]
    missing_appendices = []
    for letter in appendix_letters:
        if not re.search(rf'^##\s+PHỤ\s+LỤC\s+{letter}\b', md_raw, re.MULTILINE | re.IGNORECASE):
            missing_appendices.append(letter)
    print(f"Missing Appendix headings: {missing_appendices}")

    # 6. Table Count & Anchors
    table_headers = re.findall(r'^(?:<a id="bang-([^"]+)"></a>\s*\n)?###\s+Bảng\s+([A-Za-z0-9\._]+)\s*[-–—:]\s*(.*)', md_raw, re.MULTILINE)
    table_pipe_count = len(re.findall(r'^\|[-|\s]+\|$', md_raw, re.MULTILINE))
    print(f"Markdown Technical Table Headers matched: {len(table_headers)}")
    print(f"Markdown Pipe Separator rows (`|---|`): {table_pipe_count}")

    # Save results to JSON
    results = {
        "total_docx_paragraphs": len(doc.paragraphs),
        "total_docx_non_empty": len(all_docx_paras),
        "md_lines": len(md_lines),
        "md_chars": len(md_raw),
        "exact_full_matches": len(matched_full),
        "exact_punct_matches": len(matched_punct),
        "fuzzy_matches": len(matched_fuzzy),
        "unmatched_count": len(unmatched),
        "unmatched_items": unmatched,
        "fuzzy_items": matched_fuzzy,
        "sec_1_4_definitions_count": len(sec_1_4_headings),
        "sec_1_4_missing": missing_defs,
        "sec_1_4_dot_split_bugs": len(sec_1_4_dot_split),
        "chapter_4_space_headings": len(space_h4),
        "chapter_7_space_headings": len(space_h7),
        "missing_subdots": missing_subdots,
        "fake_table_headings": len(fake_table_headings),
        "heading_distribution": h_dist,
        "missing_appendices": missing_appendices,
        "table_headers_count": len(table_headers)
    }

    OUT_JSON.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\nSaved detailed JSON to {OUT_JSON}")

if __name__ == "__main__":
    run_suite()
