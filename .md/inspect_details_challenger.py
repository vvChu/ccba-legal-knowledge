"""
Detailed Inspection of Specific Anomaly & Subdot Findings
Author: challenger_m1_1
"""
import sys
import re
import json
from pathlib import Path
import docx

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

MD_PATH = Path(r"d:\GitHubProjects\ccba-legal-knowledge\legal_docs\02_qcvn\qcvn_06_2022_bxd\qcvn_06_2022_bxd.md")
DOCX_PATH = Path(r"d:\GitHubProjects\ccba-legal-knowledge\.md\extracted_docs\qcvn_06_2022_bxd\qcvn_06_2022_bxd.docx")

md_text = MD_PATH.read_text(encoding="utf-8")
doc = docx.Document(str(DOCX_PATH))

print("=== CHECKING SUBDOTS IN MARKDOWN ===")
subdots = [
    "2.1.1.1", "2.1.1.2", "2.2.1.1", "2.2.1.2", "2.2.1.3", "2.5.6.3.3",
    "5.1.1.1", "5.1.1.2", "5.1.1.3", "5.1.1.4",
    "6.2.1.1", "6.2.1.2", "6.2.1.3", "6.2.1.4"
]

for sd in subdots:
    matches = re.findall(rf'.{{0,40}}{re.escape(sd)}.{{0,60}}', md_text)
    print(f"\nSubdot: {sd} (Matches found: {len(matches)})")
    for m in matches[:3]:
        print(f"  Snippet: {repr(m)}")

print("\n=== CHECKING FRONTMATTER PARAGRAPHS (DOCX 0 to 22) ===")
for i in range(23):
    p = doc.paragraphs[i]
    print(f"Docx P#{i:2d}: {repr(p.text)}")

print("\n=== CHECKING 13 PUNCTUATION-STRIPPED MATCHES ===")
# Load previous results
with open(r'd:\GitHubProjects\ccba-legal-knowledge\.md\challenger_m1_adversarial_results.json', 'r', encoding='utf-8') as f:
    res = json.load(f)

# Let's find exactly which 13 paragraphs were matched_punct
clean_md = md_text.replace('\xa0', ' ').replace('\u200b', '')
clean_md_norm = re.sub(r'<a id="[^"]+"></a>', '', clean_md)
clean_md_norm = re.sub(r'[#\*_`]', '', clean_md_norm)
clean_md_norm = re.sub(r'\s+', ' ', clean_md_norm).strip()

def strip_punct(t):
    return re.sub(r'[^\w\d]+', '', t.lower())

punct_md = strip_punct(clean_md)

punct_items = []
for idx, p in enumerate(doc.paragraphs):
    raw = p.text.strip()
    if not raw:
        continue
    norm = re.sub(r'\s+', ' ', raw.replace('\xa0', ' ')).strip()
    p_strip = strip_punct(raw)
    if norm not in clean_md_norm:
        if p_strip in punct_md:
            punct_items.append((idx, raw, norm))

print(f"Total punct items found: {len(punct_items)}")
for idx, raw, norm in punct_items:
    print(f"\nDocx P#{idx}:")
    print(f"  DOCX RAW: {repr(raw)}")
    # Find matching snippet in MD
    # Look for close substring
    words = norm.split()[:6]
    search_snip = " ".join(words)
    loc = clean_md_norm.find(search_snip)
    if loc != -1:
        print(f"  MD MATCH: {repr(clean_md_norm[loc:loc+len(norm)+20])}")
    else:
        print("  MD MATCH: (exact phrase not found, punct matched)")

