import os
import re
import docx

DOCX_PATH = r"d:\GitHubProjects\ccba-legal-knowledge\.md\extracted_docs\qcvn_06_2022_bxd\qcvn_06_2022_bxd.docx"
MD_PATH = r"d:\GitHubProjects\ccba-legal-knowledge\legal_docs\02_qcvn\qcvn_06_2022_bxd\qcvn_06_2022_bxd.md"

def test_docx_structure():
    doc = docx.Document(DOCX_PATH)
    print(f"Total docx paragraphs: {len(doc.paragraphs)}")
    print(f"Total docx tables: {len(doc.tables)}")
    
    # Non-empty paragraphs
    non_empty = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    print(f"Total non-empty paragraphs: {len(non_empty)}")

def test_md_headings_and_anomalies():
    with open(MD_PATH, "r", encoding="utf-8") as f:
        md_text = f.read()
        md_lines = md_text.splitlines()

    print(f"Total MD lines: {len(md_lines)}")
    print(f"Total MD length: {len(md_text)} chars")
    
    # 1. Check dot-split in Section 1.4: 1.4.X.Y
    dot_split_1_4 = re.findall(r"^#{1,6}\s+1\.4\.\d+\.\d+.*", md_text, re.MULTILINE)
    print(f"[TEST 1] Dot-split in 1.4 (e.g. 1.4.1.7): {len(dot_split_1_4)} found -> {dot_split_1_4[:5]}")

    # 2. Check space-separated clause headers: ### X Y
    space_headings = re.findall(r"^#{1,6}\s+\d+\s+\d+.*", md_text, re.MULTILINE)
    print(f"[TEST 2] Space-separated headings (e.g. ### 4 1, ### 24 0): {len(space_headings)} found -> {space_headings[:5]}")
    
    # 3. Check appendix space-separated headings: ### A 1
    appendix_space = re.findall(r"^#{1,6}\s+[A-I]\s+\d+.*", md_text, re.MULTILINE)
    print(f"[TEST 3] Appendix space headings (e.g. ### A 1): {len(appendix_space)} found -> {appendix_space[:5]}")

    # 4. Check missing dots in 4-level or 5-level clauses
    missing_subdots = re.findall(r"^#{1,6}\s+(?:2\.1\.11|2\.1\.12|2\.2\.11|2\.2\.12|2\.2\.13|2\.5\.6\.33|5\.1\.11|5\.1\.12|5\.1\.13|5\.1\.14|6\.2\.11|6\.2\.12|6\.2\.13|6\.2\.14)\b.*", md_text, re.MULTILINE)
    print(f"[TEST 4] Specific missing subdots (e.g. 2.1.11): {len(missing_subdots)} found -> {missing_subdots[:5]}")

    # 5. Check all headings in MD
    headings = [line for line in md_lines if line.startswith("#")]
    print(f"Total headings: {len(headings)}")
    
    h_counts = {}
    for h in headings:
        level = len(h) - len(h.lstrip("#"))
        h_counts[level] = h_counts.get(level, 0) + 1
    print(f"Heading counts by level: {sorted(h_counts.items())}")

    # 6. Check Section 1.4 definitions (1.4.1 to 1.4.72)
    defs_found = {}
    for line in md_lines:
        m = re.match(r"^#{1,6}\s+1\.4\.(\d+)\b", line)
        if m:
            num = int(m.group(1))
            defs_found[num] = line
            
    print(f"Total 1.4 definitions found: {len(defs_found)} / 72")
    missing_defs = [i for i in range(1, 73) if i not in defs_found]
    if missing_defs:
        print(f"Missing 1.4 definitions: {missing_defs}")
    else:
        print("All 1.4 definitions from 1.4.1 to 1.4.72 are present!")

if __name__ == "__main__":
    test_docx_structure()
    print("=" * 60)
    test_md_headings_and_anomalies()
