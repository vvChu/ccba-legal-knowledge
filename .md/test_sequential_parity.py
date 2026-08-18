"""
Strict Sequential Parity and Text Equality Harness
Author: challenger_m1_1
"""
import sys
import re
import json
import docx
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

DOCX_PATH = Path(r"d:\GitHubProjects\ccba-legal-knowledge\.md\extracted_docs\qcvn_06_2022_bxd\qcvn_06_2022_bxd.docx")
MD_PATH = Path(r"d:\GitHubProjects\ccba-legal-knowledge\legal_docs\02_qcvn\qcvn_06_2022_bxd\qcvn_06_2022_bxd.md")

doc = docx.Document(str(DOCX_PATH))
md_raw = MD_PATH.read_text(encoding="utf-8")

def clean_for_diff(t):
    t = t.replace('\xa0', ' ').replace('\u200b', '')
    t = re.sub(r'<a id="[^"]+"></a>', '', t)
    t = re.sub(r'[#\*_`]', '', t)
    t = re.sub(r'\s+', ' ', t).strip()
    return t

# Collect all body paragraphs from DOCX (excluding cover & TOC paragraphs 0..21)
docx_body = []
for idx in range(22, len(doc.paragraphs)):
    p = doc.paragraphs[idx]
    raw = p.text.strip()
    if raw:
        docx_body.append({
            "docx_idx": idx,
            "raw": raw,
            "cleaned": clean_for_diff(raw)
        })

print(f"Total DOCX body paragraphs to audit (from 'Lời nói đầu' to end): {len(docx_body)}")

# Clean MD text into normalized searchable structure
md_cleaned = clean_for_diff(md_raw)

# Sequential matching test
current_pos = 0
sequence_failures = []
text_diffs = []
matched_count = 0

for item in docx_body:
    target = item["cleaned"]
    pos = md_cleaned.find(target, current_pos)
    if pos != -1:
        matched_count += 1
        current_pos = pos + len(target)
    else:
        # Check if it exists earlier or later or not at all
        earlier_pos = md_cleaned.find(target)
        if earlier_pos != -1:
            sequence_failures.append({
                "docx_idx": item["docx_idx"],
                "target": target[:80],
                "reason": f"Out of sequence: found at pos {earlier_pos} but expected after {current_pos}"
            })
        else:
            # Try relaxing numbers / punctuation
            # e.g. 2.1.11 vs 2.1.1.1
            norm_target = re.sub(r'[^\w\d]+', '', target.lower())
            norm_md = re.sub(r'[^\w\d]+', '', md_cleaned.lower())
            if norm_target in norm_md:
                text_diffs.append({
                    "docx_idx": item["docx_idx"],
                    "target": target[:100],
                    "reason": "Found with minor number/punctuation variation"
                })
            else:
                text_diffs.append({
                    "docx_idx": item["docx_idx"],
                    "target": target[:100],
                    "reason": "NOT FOUND in markdown!"
                })

print("\n--- SEQUENTIAL PARITY RESULTS ---")
print(f"Strict In-Order Matches: {matched_count} / {len(docx_body)} ({matched_count/len(docx_body)*100:.2f}%)")
print(f"Sequence Out-of-Order: {len(sequence_failures)}")
print(f"Text/Number Differences: {len(text_diffs)}")

if sequence_failures:
    print("\n--- SEQUENCE FAILURES ---")
    for sf in sequence_failures[:10]:
        print(f"  P#{sf['docx_idx']}: {sf['reason']} -> {sf['target']}")

if text_diffs:
    print("\n--- TEXT / NUMBER DIFFERENCES ---")
    for td in text_diffs:
        print(f"  P#{td['docx_idx']} [{td['reason']}]: {td['target']}")

