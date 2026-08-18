import docx
import json
import re
import sys

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

docx_path = r'd:\GitHubProjects\ccba-legal-knowledge\.md\extracted_docs\qcvn_06_2022_bxd\qcvn_06_2022_bxd.docx'
md_path = r'd:\GitHubProjects\ccba-legal-knowledge\legal_docs\02_qcvn\qcvn_06_2022_bxd\qcvn_06_2022_bxd.md'

doc = docx.Document(docx_path)
print(f"Total doc.paragraphs in DOCX: {len(doc.paragraphs)}")
print(f"Total doc.tables in DOCX: {len(doc.tables)}")

all_paras = []
non_empty_count = 0
empty_count = 0

for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    is_empty = len(text) == 0
    if is_empty:
        empty_count += 1
    else:
        non_empty_count += 1
    all_paras.append({
        'index': i,
        'style': p.style.name if p.style else None,
        'text': text,
        'length': len(text),
        'is_empty': is_empty
    })

print(f"Non-empty paragraphs: {non_empty_count}")
print(f"Empty paragraphs: {empty_count}")

# Check first 25 paragraphs
print("\n--- FIRST 25 PARAGRAPHS ---")
for p in all_paras[:25]:
    if not p['is_empty']:
        print(f"p{p['index']:04d}: ({p['style']}) {p['text'][:90]}")

with open(r'd:\GitHubProjects\ccba-legal-knowledge\.md\challenger_docx_inventory.json', 'w', encoding='utf-8') as f:
    json.dump({
        'total_paragraphs': len(doc.paragraphs),
        'non_empty_count': non_empty_count,
        'empty_count': empty_count,
        'total_tables': len(doc.tables),
        'paragraphs': all_paras
    }, f, ensure_ascii=False, indent=2)

print("Saved .md/challenger_docx_inventory.json")
