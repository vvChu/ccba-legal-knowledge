import docx
import json

docx_path = r'd:\GitHubProjects\ccba-legal-knowledge\.md\extracted_docs\qcvn_06_2022_bxd\qcvn_06_2022_bxd.docx'
md_path = r'd:\GitHubProjects\ccba-legal-knowledge\legal_docs\02_qcvn\qcvn_06_2022_bxd\qcvn_06_2022_bxd.md'

doc = docx.Document(docx_path)
with open(md_path, 'r', encoding='utf-8') as f:
    md_content = f.read()

# Let's inspect each Appendix in DOCX:
# Phụ lục A: p1212 to 1466
# Phụ lục B: p1467 to 1496
# Phụ lục C: p1497 to 1570
# Phụ lục D: p1571 to 1738
# Phụ lục E: p1739 to 1768
# Phụ lục F: p1769 to 1796
# Phụ lục G: p1797 to 1849
# Phụ lục H: p1850 to 1967
# Phụ lục I: p1968 to end

appendix_ranges = [
    ("A", 1212, 1466),
    ("B", 1467, 1496),
    ("C", 1497, 1570),
    ("D", 1571, 1738),
    ("E", 1739, 1768),
    ("F", 1769, 1796),
    ("G", 1797, 1849),
    ("H", 1850, 1967),
    ("I", 1968, len(doc.paragraphs)-1)
]

appendix_audit = {}

for name, start_p, end_p in appendix_ranges:
    total_p = 0
    found_p = 0
    missing_samples = []
    
    for p_idx in range(start_p, min(end_p + 1, len(doc.paragraphs))):
        t = doc.paragraphs[p_idx].text.strip()
        if not t:
            continue
        total_p += 1
        
        # Check if first 30 chars or full text is in MD
        search_snippet = t[:30].strip()
        if len(search_snippet) >= 10 and search_snippet in md_content:
            found_p += 1
        elif len(t) < 10 and t in md_content:
            found_p += 1
        else:
            missing_samples.append({
                'p_idx': p_idx,
                'text': t[:100]
            })
            
    appendix_audit[name] = {
        'range': f"p{start_p}-p{end_p}",
        'non_empty_paragraphs': total_p,
        'matched_in_md': found_p,
        'match_rate': f"{(found_p/total_p*100):.1f}%" if total_p > 0 else "0%",
        'missing_count': len(missing_samples),
        'missing_samples': missing_samples[:10]
    }

with open(r'd:\GitHubProjects\ccba-legal-knowledge\.md\appendix_audit_match.json', 'w', encoding='utf-8') as out:
    json.dump(appendix_audit, out, ensure_ascii=False, indent=2)

print("Appendix audit match completed.")
