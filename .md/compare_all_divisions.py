import docx
import json
import re

docx_path = r'd:\GitHubProjects\ccba-legal-knowledge\.md\extracted_docs\qcvn_06_2022_bxd\qcvn_06_2022_bxd.docx'
extracted_md_path = r'd:\GitHubProjects\ccba-legal-knowledge\.md\extracted_docs\qcvn_06_2022_bxd\qcvn_06_2022_bxd.md'
target_md_path = r'd:\GitHubProjects\ccba-legal-knowledge\legal_docs\02_qcvn\qcvn_06_2022_bxd\qcvn_06_2022_bxd.md'

doc = docx.Document(docx_path)
with open(extracted_md_path, 'r', encoding='utf-8') as f:
    ext_md = f.read()

with open(target_md_path, 'r', encoding='utf-8') as f:
    tgt_md = f.read()

print(f"DOCX paragraphs: {len(doc.paragraphs)}")
print(f"Extracted MD length: {len(ext_md)}")
print(f"Target MD length: {len(tgt_md)}")

# Check what each file contains for Chapters 1-7 and Appendices A-I
divisions = [
    ("CH1", "1. QUY ĐỊNH CHUNG", 5, 233),
    ("CH2", "2. PHÂN LOẠI KỸ THUẬT VỀ CHÁY", 234, 442),
    ("CH3", "3. BẢO ĐẢM AN TOÀN CHO NGƯỜI", 443, 715),
    ("CH4", "4. NGĂN CHẶN CHÁY LAN", 716, 829),
    ("CH5", "5. CẤP NƯỚC CHỮA CHÁY", 830, 1079),
    ("CH6", "6. CHỮA CHÁY VÀ CỨU NẠN", 1080, 1205),
    ("CH7", "7. TỔ CHỨC THỰC HIỆN", 1206, 1211),
    ("PLA", "PHỤ LỤC A", 1212, 1466),
    ("PLB", "PHỤ LỤC B", 1467, 1496),
    ("PLC", "PHỤ LỤC C", 1497, 1570),
    ("PLD", "PHỤ LỤC D", 1571, 1738),
    ("PLE", "PHỤ LỤC E", 1739, 1768),
    ("PLF", "PHỤ LỤC F", 1769, 1796),
    ("PLG", "PHỤ LỤC G", 1797, 1849),
    ("PLH", "PHỤ LỤC H", 1850, 1967),
    ("PLI", "PHỤ LỤC I", 1968, len(doc.paragraphs)-1),
]

comparison = []
for code, title, start_p, end_p in divisions:
    p_texts = [doc.paragraphs[i].text.strip() for i in range(start_p, min(end_p+1, len(doc.paragraphs))) if doc.paragraphs[i].text.strip()]
    
    # Check match in Target MD
    tgt_matches = 0
    ext_matches = 0
    for t in p_texts:
        snip = t[:30].strip()
        if len(snip) >= 8:
            if snip in tgt_md:
                tgt_matches += 1
            if snip in ext_md:
                ext_matches += 1
                
    total = len(p_texts)
    comparison.append({
        'code': code,
        'title': title,
        'docx_para_count': total,
        'target_md_matches': tgt_matches,
        'target_md_rate': f"{(tgt_matches/total*100):.1f}%" if total else "0%",
        'ext_md_matches': ext_matches,
        'ext_md_rate': f"{(ext_matches/total*100):.1f}%" if total else "0%"
    })

with open(r'd:\GitHubProjects\ccba-legal-knowledge\.md\comprehensive_division_audit.json', 'w', encoding='utf-8') as out:
    json.dump(comparison, out, ensure_ascii=False, indent=2)

print("Comprehensive division audit saved.")
