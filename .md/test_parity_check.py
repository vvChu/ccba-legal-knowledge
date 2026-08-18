import docx
import json
import re
import sys

sys.stdout.reconfigure(encoding="utf-8")

docx_path = r'd:\GitHubProjects\ccba-legal-knowledge\.md\extracted_docs\qcvn_06_2022_bxd\qcvn_06_2022_bxd.docx'
md_path = r'd:\GitHubProjects\ccba-legal-knowledge\.md\test_reconstructed_qcvn06.md'

doc = docx.Document(docx_path)
with open(md_path, 'r', encoding='utf-8') as f:
    md_content = f.read()

clean_md = re.sub(r'[#\*_`\<\>\[\]\(\)]', '', md_content)
clean_md = ' '.join(clean_md.split())

def norm_clean(t):
    return re.sub(r'[^\w\d]+', '', t.lower())

norm_md_str = norm_clean(clean_md)

section_boundaries = [
    ("Chapter 1", 22, 233),
    ("Chapter 2", 234, 442),
    ("Chapter 3", 443, 715),
    ("Chapter 4", 716, 829),
    ("Chapter 5", 830, 1079),
    ("Chapter 6", 1080, 1205),
    ("Chapter 7", 1206, 1211),
    ("Appendix A", 1212, 1466),
    ("Appendix B", 1467, 1496),
    ("Appendix C", 1497, 1570),
    ("Appendix D", 1571, 1738),
    ("Appendix E", 1739, 1768),
    ("Appendix F", 1769, 1796),
    ("Appendix G", 1797, 1849),
    ("Appendix H", 1850, 1967),
    ("Appendix I", 1968, 2036),
]

overall_matched = 0
overall_total = 0
all_missing = []

for name, start_p, end_p in section_boundaries:
    sec_paras = []
    for i in range(start_p, min(end_p+1, len(doc.paragraphs))):
        t = doc.paragraphs[i].text.strip()
        if t:
            sec_paras.append((i, t))
            
    total = len(sec_paras)
    matched = 0
    missing = []
    
    for p_idx, t in sec_paras:
        nt = norm_clean(t)
        snip = nt[:25] if len(nt) >= 25 else nt
        if len(snip) >= 5 and snip in norm_md_str:
            matched += 1
        elif len(snip) < 5 and (t.lower() in clean_md.lower() or nt in norm_md_str):
            matched += 1
        else:
            missing.append({'docx_p_idx': p_idx, 'text': t})
            all_missing.append({'section': name, 'docx_p_idx': p_idx, 'text': t})
            
    overall_total += total
    overall_matched += matched
    
    rate = f"{(matched/total*100):.2f}%" if total else "0%"
    print(f"{name:12s}: {matched:3d}/{total:3d} ({rate}) - Missing: {len(missing)}")
    if missing:
        for m in missing[:5]:
            print(f"   -> p[{m['docx_p_idx']}]: {m['text'][:80]}")

print("-----------------------------------------------------------------")
print(f"OVERALL PARITY: {overall_matched}/{overall_total} ({(overall_matched/overall_total*100):.2f}%)")
print(f"TOTAL MISSING : {len(all_missing)}")
