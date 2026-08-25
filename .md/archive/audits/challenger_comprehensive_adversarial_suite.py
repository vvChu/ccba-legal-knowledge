import docx
import json
import re
import sys
import unicodedata
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

docx_path = r'd:\GitHubProjects\ccba-legal-knowledge\.md\extracted_docs\qcvn_06_2022_bxd\qcvn_06_2022_bxd.docx'
md_path = r'd:\GitHubProjects\ccba-legal-knowledge\legal_docs\02_qcvn\qcvn_06_2022_bxd\qcvn_06_2022_bxd.md'

doc = docx.Document(docx_path)
with open(md_path, 'r', encoding='utf-8') as f:
    raw_md = f.read()

lines = raw_md.splitlines()

print("=================================================================")
print("  ADVERSARIAL STRESS TEST SUITE — CHALLENGER M1_2_1              ")
print("=================================================================")

# -------------------------------------------------------------
# TEST 1: Headings & Hierarchy Invariants
# -------------------------------------------------------------
print("\n--- TEST 1: Headings & Hierarchy Invariants ---")
h1_lines = [(i+1, l) for i, l in enumerate(lines) if re.match(r'^#\s+', l)]
h2_lines = [(i+1, l) for i, l in enumerate(lines) if re.match(r'^##\s+', l)]
h3_lines = [(i+1, l) for i, l in enumerate(lines) if re.match(r'^###\s+', l)]
h4_lines = [(i+1, l) for i, l in enumerate(lines) if re.match(r'^####\s+', l)]
h5_lines = [(i+1, l) for i, l in enumerate(lines) if re.match(r'^#####\s+', l)]
h6_lines = [(i+1, l) for i, l in enumerate(lines) if re.match(r'^######\s+', l)]

print(f"H1 count: {len(h1_lines)} (Expected: exactly 1)")
for ln, l in h1_lines:
    print(f"  Line {ln}: {l}")

print(f"H2 count: {len(h2_lines)} (Chapters, Appendices, Major Titles)")
print(f"H3 count: {len(h3_lines)} (Sections)")
print(f"H4 count: {len(h4_lines)} (Clauses)")
print(f"H5 count: {len(h5_lines)} (Sub-clauses)")
print(f"H6 count: {len(h6_lines)}")

assert len(h1_lines) == 1, f"Expected 1 H1, found {len(h1_lines)}"

# -------------------------------------------------------------
# TEST 2: Dot Split, Missing Subdot, and Space-Heading Anomalies
# -------------------------------------------------------------
print("\n--- TEST 2: Heading Numbering & Subdot Anomalies ---")
# 1. Check space headings: e.g. "### 4 1" or "#### 1 4"
space_headings = []
for i, l in enumerate(lines):
    m = re.match(r'^(#{1,6})\s+(\d+)\s+(\d+)\b', l)
    if m:
        space_headings.append((i+1, l))

print(f"Space headings found: {len(space_headings)}")
for ln, l in space_headings:
    print(f"  Line {ln}: {l}")

# 2. Check dot-split anomalies in Section 1.4: e.g. 1.4.1.7 to 1.4.7.2
dot_split_anomalies = []
for i, l in enumerate(lines):
    m = re.search(r'^(#{1,6})\s+1\.4\.\d+\.\d+\b', l)
    if m:
        dot_split_anomalies.append((i+1, l))

print(f"Dot-split anomalies in Section 1.4: {len(dot_split_anomalies)}")
for ln, l in dot_split_anomalies:
    print(f"  Line {ln}: {l}")

# 3. Check known missing subdots: 2.1.11, 2.1.12, 2.2.11, 2.2.12, 2.2.13, 5.1.11..14, 6.2.11..14
corrupted_subdots = []
target_subdot_regex = re.compile(r'^(#{1,6})\s+(2\.1\.1[12]|2\.2\.1[123]|5\.1\.1[1-4]|6\.2\.1[1-4])\b')
for i, l in enumerate(lines):
    if target_subdot_regex.search(l):
        corrupted_subdots.append((i+1, l))

print(f"Corrupted subdots found: {len(corrupted_subdots)}")
for ln, l in corrupted_subdots:
    print(f"  Line {ln}: {l}")

# 4. Verify normalized subdots are present
normalized_subdots = []
valid_subdot_regex = re.compile(r'^(#{1,6})\s+(2\.1\.1\.[12]|2\.2\.1\.[123]|2\.2\.2\.1|5\.1\.1\.[1-4]|6\.2\.1\.[1-4])\b')
for i, l in enumerate(lines):
    m = valid_subdot_regex.search(l)
    if m:
        normalized_subdots.append((i+1, l.strip()))

print(f"Normalized 4-level subdots present: {len(normalized_subdots)} (Expected: 14)")
for ln, l in normalized_subdots:
    print(f"  Line {ln}: {l}")

# -------------------------------------------------------------
# TEST 3: Anchor Integrity & Correspondence
# -------------------------------------------------------------
print("\n--- TEST 3: Anchor Integrity & Correspondence ---")
anchor_regex = re.compile(r'<a id="([^"]+)"></a>')
all_anchors = []
for i, l in enumerate(lines):
    for m in anchor_regex.finditer(l):
        all_anchors.append((i+1, m.group(1), l))

print(f"Total HTML anchor tags: {len(all_anchors)}")
anchor_ids = [a[1] for a in all_anchors]
duplicate_anchors = set([x for x in anchor_ids if anchor_ids.count(x) > 1])
print(f"Duplicate anchors: {len(duplicate_anchors)}")
if duplicate_anchors:
    print(f"  Duplicates: {duplicate_anchors}")

# Check anchor format
malformed_anchors = [a for a in all_anchors if not re.match(r'^(chuong-\d+|phu-luc-[a-z]|muc-[a-z0-9-]+|bang-[a-z0-9-]+)$', a[1])]
print(f"Non-standard / Malformed anchor formats: {len(malformed_anchors)}")
for ln, aid, l in malformed_anchors[:5]:
    print(f"  Line {ln}: {aid} in '{l}'")

# -------------------------------------------------------------
# TEST 4: Special Characters, Encoding & Formula Symbols
# -------------------------------------------------------------
print("\n--- TEST 4: Special Characters, Encoding & Formula Symbols ---")
null_bytes = raw_md.count('\x00')
replacement_chars = raw_md.count('\ufffd')
print(f"Null bytes: {null_bytes}")
print(f"Unicode replacement characters (\\ufffd): {replacement_chars}")

# Check key Greek letters and math symbols
math_symbols = ['Δ', 'ρ', 'μ', '≤', '≥', '±', '×', '°C', 'm²', 'm³', 'W/m²', 'kW/m²']
for sym in math_symbols:
    count_docx = sum(p.text.count(sym) for p in doc.paragraphs)
    count_md = raw_md.count(sym)
    print(f"Symbol '{sym}': DOCX={count_docx}, MD={count_md}")

# -------------------------------------------------------------
# TEST 5: Table Titles in Markdown vs DOCX
# -------------------------------------------------------------
print("\n--- TEST 5: Table Titles in Markdown ---")
docx_table_titles = []
for p in doc.paragraphs:
    t = p.text.strip()
    if re.match(r'^Bảng\s+[A-Z0-9\.]+', t, re.IGNORECASE):
        docx_table_titles.append(t)

print(f"DOCX table title paragraphs: {len(docx_table_titles)}")
missing_table_titles = []
for tt in docx_table_titles:
    # normalize
    norm_tt = re.sub(r'\s+', ' ', tt).strip()
    # check if table title appears in markdown
    if norm_tt.lower() not in raw_md.lower() and re.sub(r'[\s\-]+', '', norm_tt).lower() not in re.sub(r'[\s\-]+', '', raw_md).lower():
        missing_table_titles.append(tt)

print(f"Missing table titles in MD: {len(missing_table_titles)}")
for mtt in missing_table_titles:
    print(f"  Missing: {mtt}")

# -------------------------------------------------------------
# TEST 6: Sequential Clause Order Verification
# -------------------------------------------------------------
print("\n--- TEST 6: Sequential Clause Order ---")
# Extract all clause numbers from headings in MD
md_clause_nums = []
for i, l in enumerate(lines):
    m = re.match(r'^(#{2,6})\s+([A-Z0-9]+(?:\.[A-Z0-9]+)+)\b', l)
    if m:
        md_clause_nums.append((i+1, m.group(1), m.group(2)))

print(f"Total numbered section/clause headings in MD: {len(md_clause_nums)}")
print(f"Sample clause headings: {md_clause_nums[:10]}")

# -------------------------------------------------------------
# Summary JSON Dump
# -------------------------------------------------------------
summary_results = {
    'h1_count': len(h1_lines),
    'h2_count': len(h2_lines),
    'h3_count': len(h3_lines),
    'h4_count': len(h4_lines),
    'h5_count': len(h5_lines),
    'space_headings_count': len(space_headings),
    'dot_split_anomalies_count': len(dot_split_anomalies),
    'corrupted_subdots_count': len(corrupted_subdots),
    'normalized_subdots_count': len(normalized_subdots),
    'total_anchors': len(all_anchors),
    'duplicate_anchors': list(duplicate_anchors),
    'null_bytes': null_bytes,
    'replacement_chars': replacement_chars,
    'docx_table_titles_count': len(docx_table_titles),
    'missing_table_titles_count': len(missing_table_titles)
}

with open(r'd:\GitHubProjects\ccba-legal-knowledge\.md\challenger_comprehensive_adversarial_results.json', 'w', encoding='utf-8') as f:
    json.dump(summary_results, f, ensure_ascii=False, indent=2)

print("\n=================================================================")
print("ADVERSARIAL STRESS SUITE FINISHED.")
print("=================================================================")
