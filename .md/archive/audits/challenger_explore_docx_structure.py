import docx
import json
import re

docx_path = r'd:\GitHubProjects\ccba-legal-knowledge\.md\extracted_docs\qcvn_06_2022_bxd\qcvn_06_2022_bxd.docx'
md_path = r'd:\GitHubProjects\ccba-legal-knowledge\legal_docs\02_qcvn\qcvn_06_2022_bxd\qcvn_06_2022_bxd.md'

doc = docx.Document(docx_path)
print(f"Total doc.paragraphs in DOCX: {len(doc.paragraphs)}")
print(f"Total doc.tables in DOCX: {len(doc.tables)}")

# Inspect first 30 paragraphs
print("\n--- FIRST 30 PARAGRAPHS IN DOCX ---")
for i in range(min(30, len(doc.paragraphs))):
    text = doc.paragraphs[i].text.strip()
    print(f"p{i:04d}: [{doc.paragraphs[i].style.name if doc.paragraphs[i].style else 'NoStyle'}] {text[:80]}")

# Inspect last 20 paragraphs
print("\n--- LAST 20 PARAGRAPHS IN DOCX ---")
for i in range(max(0, len(doc.paragraphs)-20), len(doc.paragraphs)):
    text = doc.paragraphs[i].text.strip()
    print(f"p{i:04d}: [{doc.paragraphs[i].style.name if doc.paragraphs[i].style else 'NoStyle'}] {text[:80]}")
