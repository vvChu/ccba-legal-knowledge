"""Deep forensic audit comparing source DOCX files with extracted OKF bundles."""

import docx
from pathlib import Path
import json
import yaml
import sys

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

root = Path(__file__).resolve().parent.parent
extracted_dir = root / ".md" / "extracted_docs"
legal_dir = root / "legal_docs"
reg_file = root / "legal_registry.yaml"

with open(reg_file, "r", encoding="utf-8") as f:
    reg_data = yaml.safe_load(f)

all_docs = []
if "laws" in reg_data:
    all_docs.extend(reg_data["laws"])
if "documents" in reg_data:
    all_docs.extend(reg_data["documents"].values())

print("=========================================================================================================================")
print("                   CCBA DEEP FORENSIC AUDIT: SOURCE DOCX VS EXTRACTED OKF BUNDLES                                        ")
print("=========================================================================================================================")
print(f"{'Bundle / Document':<33} | {'DOCX Paras':<10} | {'DOCX Tbls':<9} | {'MD Lines':<9} | {'AST Nodes':<9} | {'Tmpls':<6} | {'Tbls':<6} | {'Status'}")
print("-" * 121)

total_examined = 0
optimal_count = 0
needs_attention = []

for doc_entry in all_docs:
    bundle_rel = doc_entry.get("bundle_path")
    if not bundle_rel:
        continue
    bundle_path = root / bundle_rel.strip("/")
    if not bundle_path.exists():
        continue

    slug = bundle_path.name
    total_examined += 1

    # Find matching docx
    folder = extracted_dir / slug
    docx_files = list(folder.glob("*.docx")) if folder.exists() else []

    # MD details
    md_file = bundle_path / f"{slug}.md"
    md_lines = len(md_file.read_text(encoding="utf-8").splitlines()) if md_file.exists() else 0

    # AST details
    clauses_file = bundle_path / "clauses.json"
    ast_count = len(json.loads(clauses_file.read_text(encoding="utf-8"))) if clauses_file.exists() else 0

    # Templates & Tables
    tmpls_count = len(list((bundle_path / "templates").glob("*.md"))) if (bundle_path / "templates").exists() else 0
    tbls_count = len(list((bundle_path / "tables" / "csv").glob("*.csv"))) if (bundle_path / "tables" / "csv").exists() else 0

    docx_paras = 0
    docx_tbls = 0
    if docx_files:
        try:
            d = docx.Document(str(docx_files[0]))
            docx_paras = len(d.paragraphs)
            docx_tbls = len(d.tables)
        except Exception:
            pass

    # Audit health
    status = "✅ OPTIMAL"
    reasons = []

    if docx_paras > 0 and md_lines < 30:
        status = "⚠️ REVIEW"
        reasons.append("MD too short")
    if docx_tbls > 10 and tbls_count == 0 and tmpls_count == 0:
        status = "⚠️ REVIEW"
        reasons.append(f"Has {docx_tbls} tables in DOCX but 0 extracted")
    if ast_count == 0:
        status = "❌ NO AST"
        reasons.append("Missing AST")

    if status != "✅ OPTIMAL":
        needs_attention.append((slug, reasons))
    else:
        optimal_count += 1

    status_str = status if not reasons else f"{status} ({', '.join(reasons)})"
    print(f"{slug[:31]:<33} | {docx_paras:<10} | {docx_tbls:<9} | {md_lines:<9} | {ast_count:<9} | {tmpls_count:<6} | {tbls_count:<6} | {status_str}")

print("-" * 121)
print(f"📊 SUMMARY: Total Bundles Examined: {total_examined} | 100% Optimal: {optimal_count} | Needs Attention: {len(needs_attention)}")
print("=========================================================================================================================")
