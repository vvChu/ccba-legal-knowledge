import sys
import docx
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
from docx.table import Table
import re

sys.stdout.reconfigure(encoding="utf-8")

docx_path = r"d:\GitHubProjects\ccba-legal-knowledge\.md\extracted_docs\qcvn_06_2022_bxd\qcvn_06_2022_bxd.docx"
doc = docx.Document(docx_path)

def table_to_gfm(tbl: Table) -> str:
    """Convert python-docx Table into clean GFM pipe table string."""
    if not tbl.rows:
        return ""
    
    # Extract matrix of cell texts
    matrix = []
    for row in tbl.rows:
        row_cells = []
        for cell in row.cells:
            # Clean text in cell
            cell_txt = " ".join(cell.text.split())
            # Replace pipes to avoid breaking GFM
            cell_txt = cell_txt.replace("|", "\\|")
            row_cells.append(cell_txt)
        matrix.append(row_cells)
    
    # Dedup adjacent identical cells caused by cell merging in python-docx
    # In python-docx, merged cells in a row return the same cell object / text
    cleaned_matrix = []
    for r_idx, row in enumerate(tbl.rows):
        cleaned_row = []
        seen_cells = set()
        for c_idx, cell in enumerate(row.cells):
            cell_elem_id = id(cell._tc)
            if cell_elem_id not in seen_cells:
                seen_cells.add(cell_elem_id)
                txt = " ".join(cell.text.split()).replace("|", "\\|")
                cleaned_row.append(txt)
        cleaned_matrix.append(cleaned_row)
        
    # Standardize column count across rows
    max_cols = max(len(r) for r in cleaned_matrix) if cleaned_matrix else 0
    if max_cols == 0:
        return ""
        
    for r in cleaned_matrix:
        if len(r) < max_cols:
            r.extend([""] * (max_cols - len(r)))
            
    header_row = cleaned_matrix[0]
    delimiter_row = ["---"] * max_cols
    data_rows = cleaned_matrix[1:]
    
    lines = [
        "| " + " | ".join(header_row) + " |",
        "| " + " | ".join(delimiter_row) + " |"
    ]
    for r in data_rows:
        lines.append("| " + " | ".join(r) + " |")
        
    return "\n".join(lines)

print(f"Testing table extraction on all {len(doc.tables)} tables...")
for idx, tbl in enumerate(doc.tables, 1):
    gfm = table_to_gfm(tbl)
    lines = gfm.splitlines()
    print(f"Table #{idx:2d}: {len(tbl.rows)} rows -> GFM {len(lines)} lines, header cols={len(lines[0].split('|'))-2}")
