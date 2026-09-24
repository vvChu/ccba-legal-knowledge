"""CCBA Legal Knowledge Spoke: Deterministic Ground Truth Parity Engine (v2.0 Hardened).

Performs 1-to-1 deterministic parity verification between extracted OKF v2.4
knowledge (Markdown, tables/CSV, KaTeX math, figures, templates) and ground-truth
Vector PDF / DOCX assets.

Evaluates 5-dimensional Non-Compensatory Vector Parity Score (VPS):
P = <P_verbatim, P_table, P_math, P_multimodal, P_structure>

Enforces 5 Error Taxonomies & Quality Contracts for extraction skills:
1. VERBATIM_TEXT -> ccba-markdown-document-processing (Target >= 98.0%, Strict 70% Consecutive Span)
2. 2D_GRID_REGULARITY -> table-reconstructor (Target = 100.0%, Anti-Vacuous Pass)
3. KATEX_MATH -> modernize_annex_engine (Target = 100.0%)
4. MULTIMODAL_ASSET -> modernize_annex_engine (Target = 100.0%)
5. STRUCTURAL_AST -> form-template-cleaner / relative-link-patcher (Target = 100.0%)

Zero-LLM token cost. 100% deterministic on local CPU.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import time
import unicodedata
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

import docx
import fitz  # PyMuPDF
import yaml

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", line_buffering=True)

ROOT_DIR = Path(__file__).resolve().parent.parent.parent
LEGAL_DOCS_DIR = ROOT_DIR / "legal_docs"
REGISTRY_FILE = ROOT_DIR / "legal_registry.yaml"
REPORTS_DIR = ROOT_DIR / ".md" / "reports"

TARGET_VERBATIM = 98.0
TARGET_TABLE = 100.0
TARGET_MATH = 100.0
TARGET_MULTIMODAL = 100.0
TARGET_STRUCTURE = 100.0

GOLDEN_COHORTS: Dict[str, Dict[str, Any]] = {
    "cohort_1": {
        "name": "Nhóm 1: Luật & Nghị định mới",
        "description": "Cấu trúc chương/điều đồ sộ, điều khoản chuyển tiếp, hiệu lực mới.",
        "slugs": [
            "luat_xay_dung_2025_135_2025_qh15",
            "nghi_dinh_217_2026_nd_cp",
            "nghi_dinh_212_2026_nd_cp",
        ],
    },
    "cohort_2": {
        "name": "Nhóm 2: Quy chuẩn Umbrella & Bảng 2D",
        "description": "Bảng số liệu đa tầng vMerge/colMerge, chú thích chân bảng, ma trận phụ lục.",
        "slugs": [
            "qcvn_06_2022_bxd",
            "qcvn_07_2023_bxd",
        ],
    },
    "cohort_3": {
        "name": "Nhóm 3: Toán học & KaTeX Đa Dòng",
        "description": "Hàng nghìn công thức MathType OLE, phân số đa tầng, nhãn phương trình \\qquad (X).",
        "slugs": [
            "tcvn_5575_2024",
            "tcvn_9386_2025",
        ],
    },
    "cohort_4": {
        "name": "Nhóm 4: Đồ họa Đa Phương Thức & BIM",
        "description": "Sơ đồ kích thước nhân trắc học, thẻ thị giác cards, Dual-format SVG+PNG.",
        "slugs": [
            "qcvn_10_2024_bxd",
            "tcvn_iso_19650_1_2021",
        ],
    },
    "cohort_5": {
        "name": "Nhóm 5: Tiêu Chuẩn Cũ / Di Sản",
        "description": "Thuật ngữ và bảng tra cứu cũ, font mã hóa lịch sử chuyển đổi sang Unicode NFC.",
        "slugs": [
            "tcvn_3981_1985",
            "tcvn_4474_1987",
        ],
    },
}


@dataclass
class DiagnosticTicket:
    bundle_slug: str
    error_category: str
    error_code: str
    file_path: str
    location: str
    detail: str
    assigned_skill: str
    action_required: str


@dataclass
class BundleParityResult:
    slug: str
    category: str
    bundle_path: str
    p_verbatim: float = 100.0
    p_table: float = 100.0
    p_math: float = 100.0
    p_multimodal: float = 100.0
    p_structure: float = 100.0
    passed: bool = True
    details: Dict[str, Any] = field(default_factory=dict)
    tickets: List[DiagnosticTicket] = field(default_factory=list)

    @property
    def vps_tuple(self) -> Tuple[float, float, float, float, float]:
        return (
            round(self.p_verbatim, 2),
            round(self.p_table, 2),
            round(self.p_math, 2),
            round(self.p_multimodal, 2),
            round(self.p_structure, 2),
        )


def normalize_unicode_text(text: str) -> str:
    """Normalize text using Unicode NFC and standardize whitespace."""
    if not text:
        return ""
    # Enforce NFC normalization on all inputs
    text = unicodedata.normalize("NFC", text)
    text = text.replace("–", "-").replace("—", "-").replace("“", '"').replace("”", '"')
    text = text.replace("‘", "'").replace("’", "'").replace("…", "...")
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    return text.strip()


def normalize_for_matching(text: str) -> str:
    """Normalize text for strict matching, mapping KaTeX symbols to text representations."""
    text = normalize_unicode_text(text).lower()

    # Strip HTML tags
    text = re.sub(r"</?[a-zA-Z][^>]*>", " ", text)

    # Convert Markdown links to anchor text
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)

    # Filter standalone page numbers
    if re.match(r"^\d+$", text.strip()):
        return ""

    # Strip HTML entities (&nbsp;, &lt;, &gt;, &amp;, etc.)
    text = re.sub(r"&[a-zA-Z0-9#]+;", " ", text)

    # KaTeX math symbol alignment map
    greek_map = {
        r"\alpha": "α",
        r"\beta": "β",
        r"\gamma": "γ",
        r"\delta": "δ",
        r"\epsilon": "ε",
        r"\varepsilon": "ε",
        r"\eta": "η",
        r"\theta": "θ",
        r"\lambda": "λ",
        r"\mu": "μ",
        r"\nu": "ν",
        r"\xi": "ξ",
        r"\pi": "π",
        r"\rho": "ρ",
        r"\sigma": "σ",
        r"\tau": "τ",
        r"\phi": "φ",
        r"\varphi": "φ",
        r"\chi": "χ",
        r"\psi": "ψ",
        r"\omega": "ω",
    }
    for katex_sym, char in greek_map.items():
        text = text.replace(katex_sym, char)

    # Strip KaTeX formatting macro keywords
    text = re.sub(r"\\(bar|overline|vec|hat|tilde|mathbf|mathrm|text|mathit|displaystyle|limits)\b", "", text)

    # Strip math delimiters and formatting
    text = re.sub(r"[\$_\{\}\^\\]", "", text)
    text = re.sub(r"[^\w\s]", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def check_multi_span_coverage(
    p_words: list[str], norm_md: str, min_span: int = 4, min_ratio: float = 0.70
) -> bool:
    """Checks if a paragraph has cumulative consecutive span coverage >= min_ratio.

    - For paragraphs with < 5 words: entire phrase must match.
    - For paragraphs with >= 5 words: greedy non-overlapping span matching with min_span >= 4.
    This prevents false passes from generic legal boilerplate while accommodating inline formula splits.
    """
    total = len(p_words)
    if total < 5:
        return " ".join(p_words) in norm_md

    # Greedy longest span matching
    matched_indices = set()
    for span_len in range(total, min_span - 1, -1):
        for i in range(total - span_len + 1):
            if any(idx in matched_indices for idx in range(i, i + span_len)):
                continue
            span = " ".join(p_words[i : i + span_len])
            if span in norm_md:
                for idx in range(i, i + span_len):
                    matched_indices.add(idx)

    ratio = len(matched_indices) / total
    return ratio >= min_ratio


class GroundTruthParityVerifier:
    """Verifies OKF v2.4 Bundles against Vector PDF / DOCX Ground Truth."""

    def __init__(self, root_dir: Path = ROOT_DIR) -> None:
        self.root_dir = root_dir
        self.legal_docs_dir = root_dir / "legal_docs"
        self.registry = self._load_registry()

    def _load_registry(self) -> Dict[str, Any]:
        if not REGISTRY_FILE.exists():
            return {}
        try:
            return yaml.safe_load(REGISTRY_FILE.read_text(encoding="utf-8")) or {}
        except Exception:
            return {}

    def find_bundle_dir(self, slug: str) -> Optional[Path]:
        """Locate bundle folder by document slug."""
        for cat in ["01_vbpl", "02_qcvn", "03_tcvn", "04_appendices"]:
            candidate = self.legal_docs_dir / cat / slug
            if candidate.exists() and candidate.is_dir():
                return candidate
            matches = list((self.legal_docs_dir / cat).glob(f"*{slug}*"))
            if matches and matches[0].is_dir():
                return matches[0]
        return None

    def verify_bundle(self, slug: str) -> BundleParityResult:
        """Run 5-dimensional parity checks on a single document bundle."""
        bundle_dir = self.find_bundle_dir(slug)
        if not bundle_dir:
            res = BundleParityResult(
                slug=slug,
                category="unknown",
                bundle_path="",
                p_verbatim=0.0,
                passed=False,
            )
            res.tickets.append(
                DiagnosticTicket(
                    bundle_slug=slug,
                    error_category="STRUCTURAL_AST",
                    error_code="ERR_BUNDLE_NOT_FOUND",
                    file_path=str(self.legal_docs_dir),
                    location="legal_docs/",
                    detail=f"Bundle folder for slug '{slug}' not found on disk.",
                    assigned_skill="platform-loader",
                    action_required="Ensure bundle is ingested and correctly registered in legal_registry.yaml.",
                )
            )
            return res

        cat = bundle_dir.parent.name
        result = BundleParityResult(
            slug=slug,
            category=cat,
            bundle_path=str(bundle_dir.relative_to(self.root_dir)),
        )

        # Dimension 1: Verbatim Parity against Ground Truth (DOCX or Vector PDF)
        self._check_verbatim_parity(bundle_dir, result)

        # Dimension 2: 2D Grid Regularity & Footnotes (with Anti-Vacuous Pass)
        self._check_table_regularity(bundle_dir, result)

        # Dimension 3: KaTeX Math Syntax & Multiline Tags
        self._check_katex_syntax(bundle_dir, result)

        # Dimension 4: Multimodal Decoupled Assets & Cards
        self._check_multimodal_assets(bundle_dir, result)

        # Dimension 5: Structural AST & Form Templates
        self._check_structural_ast(bundle_dir, result)

        # Non-Compensatory Pass/Fail Gate
        passed = (
            result.p_verbatim >= TARGET_VERBATIM
            and result.p_table >= TARGET_TABLE
            and result.p_math >= TARGET_MATH
            and result.p_multimodal >= TARGET_MULTIMODAL
            and result.p_structure >= TARGET_STRUCTURE
        )
        result.passed = passed

        return result

    def _extract_all_bundle_markdown_text(self, bundle_dir: Path) -> str:
        """Extract all text recursively across bundle md files, annexes, templates, and CSVs."""
        md_text_parts: List[str] = []
        for md_file in bundle_dir.rglob("*.md"):
            if "sources" in md_file.parts or md_file.name in ("index.md", "dead_ends.md", "log.md"):
                continue
            try:
                content = md_file.read_text(encoding="utf-8")
                if content.startswith("---"):
                    parts = content.split("---", 2)
                    if len(parts) >= 3:
                        content = parts[2]
                md_text_parts.append(content)
            except Exception:
                pass

        # Also append CSV tables
        tables_dir = bundle_dir / "tables"
        if tables_dir.exists():
            for csv_f in tables_dir.rglob("*.csv"):
                try:
                    md_text_parts.append(csv_f.read_text(encoding="utf-8"))
                except Exception:
                    pass

        return "\n".join(md_text_parts)

    def _check_verbatim_parity(self, bundle_dir: Path, res: BundleParityResult) -> None:
        """Check Verbatim Text Parity using strict consecutive coverage (>= 70%)."""
        sources_dir = bundle_dir / "sources"
        docx_candidates = list(sources_dir.glob("*.docx")) if sources_dir.exists() else []
        pdf_candidates = list(sources_dir.glob("*.pdf")) if sources_dir.exists() else []
        vector_pdfs = [p for p in pdf_candidates if not p.name.endswith("_raw_scan.pdf")]

        if not docx_candidates and not vector_pdfs:
            res.p_verbatim = 100.0
            res.details["verbatim"] = {"status": "skipped", "reason": "no_source_assets"}
            return

        combined_md = self._extract_all_bundle_markdown_text(bundle_dir)
        norm_md = normalize_for_matching(combined_md)

        # 1. Prefer DOCX ground truth when present
        if docx_candidates:
            docx_path = docx_candidates[0]
            try:
                doc = docx.Document(docx_path)
                raw_paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            except Exception as exc:
                res.tickets.append(
                    DiagnosticTicket(
                        bundle_slug=bundle_dir.name,
                        error_category="VERBATIM_TEXT",
                        error_code="ERR_DOCX_READ_FAILED",
                        file_path=str(docx_path.relative_to(self.root_dir)),
                        location="sources/*.docx",
                        detail=f"Failed to read DOCX asset: {exc}",
                        assigned_skill="ccba-markdown-document-processing",
                        action_required="Ensure DOCX asset is valid.",
                    )
                )
                res.p_verbatim = 0.0
                return

            # Filter circular preamble and TOC (conforms to ADR 0021 Pure Body)
            start_idx = 0
            has_circular = any(
                k in p.upper()
                for p in raw_paras[:15]
                for k in ["THÔNG TƯ", "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", "BAN HÀNH KÈM THEO THÔNG TƯ"]
            )
            if has_circular:
                for idx, p in enumerate(raw_paras):
                    if idx == 0:
                        continue
                    if (
                        re.match(r"^(?:QCVN|TCVN)\s+[0-9]+", p.strip().upper())
                        or p.strip().upper() in ("TIÊU CHUẨN QUỐC GIA", "QUY CHUẨN KỸ THUẬT QUỐC GIA")
                        or p.strip().upper().startswith("QUY CHUẨN KỸ THUẬT QUỐC GIA")
                    ):
                        start_idx = idx
                        break

            target_paras: List[str] = []
            in_toc = False
            in_signatory = False
            in_bibliography = False
            for p_idx, p_text in enumerate(raw_paras[start_idx:], start=start_idx):
                p_strip = p_text.strip()
                p_upper = p_strip.upper()
                if p_upper in ["MỤC LỤC", "TABLE OF CONTENTS"]:
                    in_toc = True
                    continue
                if in_toc:
                    is_end = False
                    if p_strip.lower().startswith("lời nói đầu"):
                        for nxt_idx in range(p_idx + 1, min(p_idx + 5, len(raw_paras))):
                            nxt_t = raw_paras[nxt_idx].strip()
                            if nxt_t and not re.match(
                                r"^(?:Lời giới thiệu|\d+[\.\s]|Phụ lục|Thư mục)", nxt_t, re.IGNORECASE
                            ):
                                is_end = True
                                break
                    elif p_upper in ("TIÊU CHUẨN QUỐC GIA", "QUY CHUẨN KỸ THUẬT QUỐC GIA"):
                        is_end = True
                    if is_end:
                        in_toc = False

                if re.search(r"^(?:THƯ MỤC TÀI LIỆU THAM KHẢO|TÀI LIỆU THAM KHẢO)", p_strip, re.IGNORECASE):
                    in_bibliography = True
                    continue
                if in_bibliography:
                    if re.search(r"^(?:Phụ lục|PHỤ LỤC|Điều\s+\d+|CHƯƠNG|\d+[\.\s]+[A-Z])", p_strip):
                        in_bibliography = False
                    else:
                        continue

                if re.search(r"^(?:Nơi nhận:|KT\.\s*BỘ TRƯỞNG|KT\.\s*THỦ TƯỚNG|TM\.\s*CHÍNH PHỦ|THỦ TƯỚNG\b|PHÓ THỦ TƯỚNG\b|THỨ TRƯỞNG\b)", p_strip, re.IGNORECASE):
                    in_signatory = True
                    continue
                if in_signatory:
                    if re.search(r"^(?:Phụ lục|PHỤ LỤC|Điều\s+\d+|CHƯƠNG)", p_strip, re.IGNORECASE):
                        in_signatory = False
                    else:
                        continue

                if not in_toc and not in_bibliography and not in_signatory and len(p_strip.split()) >= 3:
                    target_paras.append(p_strip)

            effective_paras_count = 0
            matched_count = 0
            missing_paras = []
            for p in target_paras:
                nw = normalize_for_matching(p).split()
                if not nw:
                    continue
                if check_multi_span_coverage(nw, norm_md, min_span=4, min_ratio=0.70):
                    matched_count += 1
                    effective_paras_count += 1
                else:
                    # Check if paragraph is administrative enacting preamble (conforming to ADR 0021 Pure Body)
                    p_low = p.strip().lower()
                    if (
                        p_low.startswith("căn cứ ")
                        or p_low.startswith("theo đề nghị ")
                        or p_low.startswith("xét đề nghị ")
                        or p_low.startswith("cộng hòa xã hội chủ nghĩa việt nam")
                        or p_low.startswith("độc lập - tự do - hạnh phúc")
                        or (p_low.startswith("bộ trưởng ") and "ban hành thông tư" in p_low)
                        or p_low.startswith("chính phủ ban hành nghị định")
                        or p_low.startswith("lời nói đầu")
                        or (p_low.startswith("tcvn ") and ("thay thế " in p_low or "được xây dựng " in p_low or "biên soạn" in p_low))
                        or p_low.startswith("thư mục tài liệu tham khảo")
                        or p_low.startswith("nơi nhận:")
                        or p_low.startswith("kt. bộ trưởng")
                        or p_low.startswith("kt. thủ tướng")
                        or p_low.startswith("tm. chính phủ")
                        or p_low.startswith("thủ tướng")
                        or p_low.startswith("phó thủ tướng")
                        or p_low.startswith("thứ trưởng")
                    ):
                        continue
                    effective_paras_count += 1
                    if len(missing_paras) < 5:
                        missing_paras.append(p)

            total_paras = effective_paras_count
            rate = (matched_count / total_paras) * 100.0 if total_paras > 0 else 100.0
            res.p_verbatim = round(rate, 2)
            res.details["verbatim"] = {
                "source": "DOCX",
                "total_paragraphs": total_paras,
                "matched_paragraphs": matched_count,
                "parity_rate": round(rate, 2),
                "sample_missing": missing_paras,
            }

        # 2. Otherwise use Vector PDF
        else:
            pdf_path = vector_pdfs[0]
            try:
                doc = fitz.open(pdf_path)
            except Exception as exc:
                res.tickets.append(
                    DiagnosticTicket(
                        bundle_slug=bundle_dir.name,
                        error_category="VERBATIM_TEXT",
                        error_code="ERR_PDF_READ_FAILED",
                        file_path=str(pdf_path.relative_to(self.root_dir)),
                        location="sources/*.pdf",
                        detail=f"Failed to read Vector PDF: {exc}",
                        assigned_skill="ccba-ai-pdf-preprocessor",
                        action_required="Ensure Vector PDF is valid.",
                    )
                )
                res.p_verbatim = 0.0
                return

            # Check explicit scope from metadata.yaml (conforming to ADR 0021 / Explicit Scope Declaration)
            scope_pages = None
            meta_path = bundle_dir / "metadata.yaml"
            if meta_path.exists():
                try:
                    meta_data = yaml.safe_load(meta_path.read_text(encoding="utf-8")) or {}
                    scope_pages = meta_data.get("verification_scope", {}).get("normative_body_pages")
                except Exception:
                    pass

            if scope_pages and isinstance(scope_pages, (list, tuple)) and len(scope_pages) == 2:
                start_p, end_p = int(scope_pages[0]), int(scope_pages[1])
                page_indices = range(max(0, start_p - 1), min(len(doc), end_p))
            else:
                page_indices = range(len(doc))

            raw_blocks: List[str] = []
            for p_idx in page_indices:
                page = doc[p_idx]
                for b in page.get_text("blocks"):
                    if b[6] == 0:  # text block
                        t = b[4].strip()
                        if not t:
                            continue
                        # Filter standalone page numbers
                        if re.match(r"^\d+$", t):
                            continue
                        raw_blocks.append(t)
            doc.close()

            # Filter circular preamble and TOC (conforming to ADR 0021 Pure Body)
            start_idx = 0
            is_standard = bundle_dir.parent.name in ("02_qcvn", "03_tcvn") or bundle_dir.name.startswith(("qcvn_", "tcvn_"))
            for idx, b in enumerate(raw_blocks):
                b_s = b.strip()
                b_u = b_s.upper()
                if is_standard:
                    if (
                        re.match(r"^(?:QCVN|TCVN)\s+[0-9]+", b_u)
                        or b_u in ("TIÊU CHUẨN QUỐC GIA", "QUY CHUẨN KỸ THUẬT QUỐC GIA")
                        or b_u.startswith("QUY CHUẨN KỸ THUẬT QUỐC GIA")
                        or re.search(r"^(?:1[\.\s]+QUY ĐỊNH CHUNG|I[\.\s]+QUY ĐỊNH CHUNG)", b_s, re.IGNORECASE)
                    ):
                        start_idx = idx
                        break
                else:
                    if re.search(r"^(?:Điều\s+1\b|CHƯƠNG\s+I\b|I[\.\s]+QUY ĐỊNH CHUNG|1[\.\s]+QUY ĐỊNH CHUNG|Phần\s+1\b)", b_s, re.IGNORECASE):
                        start_idx = idx
                        break

            target_paras: List[str] = []
            in_toc = False
            in_signatory = False
            in_bibliography = False
            for b_idx, b_text in enumerate(raw_blocks[start_idx:], start=start_idx):
                b_strip = b_text.strip()
                b_upper = b_strip.upper()
                if b_upper in ["MỤC LỤC", "TABLE OF CONTENTS"]:
                    in_toc = True
                    in_bibliography = False
                    continue
                if in_toc:
                    is_end = False
                    if b_strip.lower().startswith("lời nói đầu"):
                        for nxt_idx in range(b_idx + 1, min(b_idx + 5, len(raw_blocks))):
                            nxt_t = raw_blocks[nxt_idx].strip()
                            if nxt_t and not re.match(
                                r"^(?:Lời giới thiệu|\d+[\.\s]|Phụ lục|Thư mục)", nxt_t, re.IGNORECASE
                            ):
                                is_end = True
                                break
                    elif b_upper in ("TIÊU CHUẨN QUỐC GIA", "QUY CHUẨN KỸ THUẬT QUỐC GIA"):
                        is_end = True
                    elif re.search(r"^(?:Điều\s+1\b|CHƯƠNG\s+I\b|1[\.\s]+QUY ĐỊNH CHUNG)", b_strip, re.IGNORECASE):
                        is_end = True
                    if is_end:
                        in_toc = False

                if re.search(r"^THƯ MỤC TÀI LIỆU THAM KHẢO\b", b_strip, re.IGNORECASE):
                    in_bibliography = True
                    continue
                if in_bibliography:
                    if re.search(r"^(?:Phụ lục|PHỤ LỤC|Điều\s+\d+|CHƯƠNG|\d+[\.\s]+[A-Z])", b_strip):
                        in_bibliography = False
                    else:
                        continue

                if scope_pages and re.search(r"^(?:Nơi nhận:|KT\.\s*BỘ TRƯỞNG|TM\.\s*CHÍNH PHỦ|THỦ TƯỚNG\b)", b_strip, re.IGNORECASE):
                    break

                if re.search(r"^(?:Nơi nhận:|KT\.\s*BỘ TRƯỞNG|KT\.\s*THỦ TƯỚNG|TM\.\s*CHÍNH PHỦ|THỦ TƯỚNG\b|PHÓ THỦ TƯỚNG\b|THỨ TRƯỞNG\b)", b_strip, re.IGNORECASE):
                    in_signatory = True
                    continue
                if in_signatory:
                    if re.search(r"^(?:Phụ lục|PHỤ LỤC|Điều\s+\d+|CHƯƠNG)", b_strip, re.IGNORECASE):
                        in_signatory = False
                    else:
                        continue

                if not in_toc and not in_bibliography and not in_signatory and len(b_strip.split()) >= 3:
                    target_paras.append(b_strip)

            effective_paras_count = 0
            matched_count = 0
            missing_paras = []
            for p in target_paras:
                nw = normalize_for_matching(p).split()
                if not nw:
                    continue
                if check_multi_span_coverage(nw, norm_md, min_span=4, min_ratio=0.70):
                    matched_count += 1
                    effective_paras_count += 1
                else:
                    # Check if paragraph is administrative enacting preamble / signatory (conforming to ADR 0021 Pure Body)
                    p_low = p.strip().lower()
                    if (
                        p_low.startswith("căn cứ ")
                        or p_low.startswith("theo đề nghị ")
                        or p_low.startswith("xét đề nghị ")
                        or p_low.startswith("cộng hòa xã hội chủ nghĩa việt nam")
                        or p_low.startswith("độc lập - tự do - hạnh phúc")
                        or (p_low.startswith("bộ trưởng ") and "ban hành thông tư" in p_low)
                        or p_low.startswith("chính phủ ban hành nghị định")
                        or p_low.startswith("lời nói đầu")
                        or (p_low.startswith("tcvn ") and ("thay thế " in p_low or "được xây dựng " in p_low or "biên soạn" in p_low))
                        or p_low.startswith("thư mục tài liệu tham khảo")
                        or p_low.startswith("nơi nhận:")
                        or p_low.startswith("kt. bộ trưởng")
                        or p_low.startswith("kt. thủ tướng")
                        or p_low.startswith("tm. chính phủ")
                        or p_low.startswith("thủ tướng")
                        or p_low.startswith("phó thủ tướng")
                        or p_low.startswith("thứ trưởng")
                    ):
                        continue
                    effective_paras_count += 1
                    if len(missing_paras) < 5:
                        missing_paras.append(p)

            total_paras = effective_paras_count
            rate = (matched_count / total_paras) * 100.0 if total_paras > 0 else 100.0
            res.p_verbatim = round(rate, 2)
            res.details["verbatim"] = {
                "source": "PDF_VECTOR",
                "total_paragraphs": total_paras,
                "matched_paragraphs": matched_count,
                "parity_rate": round(rate, 2),
                "sample_missing": missing_paras,
            }

        if res.p_verbatim < TARGET_VERBATIM:
            res.tickets.append(
                DiagnosticTicket(
                    bundle_slug=bundle_dir.name,
                    error_category="VERBATIM_TEXT",
                    error_code="ERR_VERBATIM_DROP",
                    file_path=str(bundle_dir.relative_to(self.root_dir)),
                    location="Normative Body Markdown",
                    detail=f"Verbatim match rate is {res.p_verbatim:.1f}% (< {TARGET_VERBATIM}%).",
                    assigned_skill="ccba-markdown-document-processing",
                    action_required="Audit missing paragraphs and ensure 1:1 verbatim reproduction.",
                )
            )

    def _check_table_regularity(self, bundle_dir: Path, res: BundleParityResult) -> None:
        """Check 2D Grid Regularity & Footnote Decoupling (with Anti-Vacuous Pass)."""
        tables_dir = bundle_dir / "tables"
        csv_dir = tables_dir / "csv"

        # Anti-Vacuous Pass: Check if DOCX source has relational data tables
        sources_dir = bundle_dir / "sources"
        docx_candidates = list(sources_dir.glob("*.docx")) if sources_dir.exists() else []
        docx_data_table_count = 0
        if docx_candidates:
            try:
                d = docx.Document(docx_candidates[0])
                for t in d.tables:
                    # Relational data tables have > 2 rows and > 1 column
                    if len(t.rows) > 2 and len(t.columns) > 1:
                        docx_data_table_count += 1
            except Exception:
                pass

        if not tables_dir.exists() or not csv_dir.exists() or not list(csv_dir.glob("*.csv")):
            if docx_data_table_count > 0:
                # Flag Vacuous Pass Error!
                res.p_table = 0.0
                res.details["table"] = {
                    "total_tables": 0,
                    "docx_tables_found": docx_data_table_count,
                    "status": "vacuous_pass_detected",
                }
                res.tickets.append(
                    DiagnosticTicket(
                        bundle_slug=bundle_dir.name,
                        error_category="2D_GRID_REGULARITY",
                        error_code="ERR_MISSING_TABLE_CSVS",
                        file_path=str(bundle_dir.relative_to(self.root_dir)),
                        location="tables/csv/",
                        detail=f"Source DOCX has {docx_data_table_count} data tables but tables/csv/ is empty.",
                        assigned_skill="table-reconstructor",
                        action_required="Extract relational data tables into tables/csv/ with valid 2D rectangular grid.",
                    )
                )
                return
            else:
                res.p_table = 100.0
                res.details["table"] = {"total_tables": 0, "status": "none"}
                return

        csv_files = sorted(list(csv_dir.glob("*.csv")))
        catalog_file = tables_dir / "tables_catalog.json"
        has_catalog = catalog_file.exists()

        valid_count = 0
        ragged_tickets = []
        footnote_tickets = []

        for csv_path in csv_files:
            rel_path = str(csv_path.relative_to(self.root_dir))
            try:
                with open(csv_path, mode="r", encoding="utf-8-sig", newline="") as f:
                    rows = list(csv.reader(f))
                while rows and not any(cell.strip() for cell in rows[-1]):
                    rows.pop()

                if not rows:
                    continue

                header_len = len(rows[0])
                if header_len == 0:
                    continue

                is_valid = True
                for r_idx, r in enumerate(rows[1:], start=2):
                    if len(r) != header_len:
                        is_valid = False
                        ragged_tickets.append((rel_path, r_idx, len(r), header_len))
                        break

                    r_text = " ".join(r).strip()
                    if re.match(r"^(?:CHÚ\s+THÍCH|CHÚ\s+DẪN|Ghi\s+chú|\(\*\))\s*:", r_text, re.IGNORECASE):
                        is_valid = False
                        footnote_tickets.append((rel_path, r_idx))
                        break

                if is_valid:
                    valid_count += 1
            except Exception:
                pass

        total = len(csv_files)
        rate = (valid_count / total) * 100.0 if total > 0 else 100.0
        res.p_table = round(rate, 2)
        res.details["table"] = {
            "total_tables": total,
            "valid_tables": valid_count,
            "has_catalog": has_catalog,
            "docx_data_tables": docx_data_table_count,
        }

        for path, r_idx, act, exp in ragged_tickets[:3]:
            res.tickets.append(
                DiagnosticTicket(
                    bundle_slug=bundle_dir.name,
                    error_category="2D_GRID_REGULARITY",
                    error_code="ERR_RAGGED_ROW",
                    file_path=path,
                    location=f"Row {r_idx}",
                    detail=f"Ragged row detected: got {act} columns, expected {exp}.",
                    assigned_skill="table-reconstructor",
                    action_required="Ensure virtual grid hierarchical forward-fill on vMerge cells.",
                )
            )

        for path, r_idx in footnote_tickets[:3]:
            res.tickets.append(
                DiagnosticTicket(
                    bundle_slug=bundle_dir.name,
                    error_category="2D_GRID_REGULARITY",
                    error_code="ERR_FOOTNOTE_IN_CSV",
                    file_path=path,
                    location=f"Row {r_idx}",
                    detail="Footnote leaked into data row.",
                    assigned_skill="table-reconstructor",
                    action_required="Decouple footnotes into tables_catalog.json metadata.",
                )
            )

    def _check_katex_syntax(self, bundle_dir: Path, res: BundleParityResult) -> None:
        """Check KaTeX Math Syntax & Multiline Tags (ADR 0038, ADR 0044)."""
        total_blocks = 0
        valid_blocks = 0
        math_tickets = []

        for md_file in bundle_dir.rglob("*.md"):
            if "sources" in md_file.parts or md_file.name in ("index.md", "dead_ends.md", "log.md"):
                continue
            try:
                content = md_file.read_text(encoding="utf-8")
            except Exception:
                continue

            if "$$" not in content:
                continue

            stripped = re.sub(r"```[\s\S]*?```", "", content)
            blocks = re.findall(r"\$\$([\s\S]*?)\$\$", stripped)
            total_blocks += len(blocks)

            for b_idx, block in enumerate(blocks, start=1):
                clean = block.strip()
                is_valid = True

                # Check 1: No HTML comments inside $$
                if "<!--" in clean and "-->" in clean:
                    is_valid = False
                    math_tickets.append(
                        (str(md_file.relative_to(self.root_dir)), b_idx, "HTML comment inside $$ block")
                    )

                # Check 2: No \tag in multiline envs
                for env in ["aligned", "cases", "gather", "matrix", "array"]:
                    if f"\\begin{{{env}}}" in clean and re.search(r"\\tag\s*\{", clean):
                        is_valid = False
                        math_tickets.append(
                            (str(md_file.relative_to(self.root_dir)), b_idx, f"\\tag used inside {env} environment")
                        )
                        break

                # Check 3: Balanced \left and \right
                lefts = len(re.findall(r"\\left[\(\[\{\.\vert]", clean))
                rights = len(re.findall(r"\\right[\)\]\}\.\vert]", clean))
                if lefts != rights:
                    is_valid = False
                    math_tickets.append(
                        (str(md_file.relative_to(self.root_dir)), b_idx, f"Unbalanced \\left ({lefts}) and \\right ({rights})")
                    )

                # Check 4: Balanced braces
                nobrace = re.sub(r"\\[\{\}]", "", clean)
                if nobrace.count("{") != nobrace.count("}"):
                    is_valid = False
                    math_tickets.append(
                        (str(md_file.relative_to(self.root_dir)), b_idx, "Unbalanced curly braces")
                    )

                if is_valid:
                    valid_blocks += 1

        rate = (valid_blocks / total_blocks) * 100.0 if total_blocks > 0 else 100.0
        res.p_math = round(rate, 2)
        res.details["math"] = {"total_blocks": total_blocks, "valid_blocks": valid_blocks}

        for path, b_idx, msg in math_tickets[:3]:
            res.tickets.append(
                DiagnosticTicket(
                    bundle_slug=bundle_dir.name,
                    error_category="KATEX_MATH",
                    error_code="ERR_KATEX_SYNTAX",
                    file_path=path,
                    location=f"Math block #{b_idx}",
                    detail=msg,
                    assigned_skill="modernize_annex_engine",
                    action_required="Fix KaTeX syntax, replace \\tag with \\qquad (X) in multiline environments.",
                )
            )

    def _check_multimodal_assets(self, bundle_dir: Path, res: BundleParityResult) -> None:
        """Check Multimodal Decoupled Assets & Cards Parity (ADR 0040)."""
        figures_dir = bundle_dir / "figures"
        if not figures_dir.exists():
            res.p_multimodal = 100.0
            res.details["multimodal"] = {"total_figures": 0, "status": "none"}
            return

        stray_binaries = []
        for ext in (".wmf", ".emf"):
            stray_binaries.extend(list(figures_dir.rglob(f"*{ext}")))

        if stray_binaries:
            for s in stray_binaries[:2]:
                res.tickets.append(
                    DiagnosticTicket(
                        bundle_slug=bundle_dir.name,
                        error_category="MULTIMODAL_ASSET",
                        error_code="ERR_STRAY_VECTOR_BINARY",
                        file_path=str(s.relative_to(self.root_dir)),
                        location="figures/",
                        detail=f"Found unviewable vector binary '{s.name}'. Must convert to SVG or PNG.",
                        assigned_skill="modernize_annex_engine",
                        action_required="Convert vector graphics to SVG/PNG dual format.",
                    )
                )

        catalog_file = figures_dir / "figures_catalog.yaml"
        cards_dir = figures_dir / "cards"

        if not catalog_file.exists():
            rate = 0.0 if stray_binaries else 100.0
            res.p_multimodal = rate
            res.details["multimodal"] = {"total_figures": 0, "has_catalog": False}
            return

        try:
            cat_data = yaml.safe_load(catalog_file.read_text(encoding="utf-8")) or {}
            fig_list = cat_data.get("figures", [])
        except Exception:
            fig_list = []

        if not fig_list:
            res.p_multimodal = 100.0
            res.details["multimodal"] = {"total_figures": 0}
            return

        valid_figs = 0
        for fig in fig_list:
            tag = str(fig.get("tag") or fig.get("id") or "").replace("hinh_", "")
            raw_slug = str(fig.get("slug") or fig.get("id") or tag).lower().replace(".", "_").replace("-", "_").replace("đ", "dd")
            f_slug = raw_slug[5:] if raw_slug.startswith("hinh_") else raw_slug

            img_rel = fig.get("image_relpath") or fig.get("path")
            img_ok = False
            if img_rel:
                img_p = bundle_dir / img_rel
                img_ok = img_p.exists() and img_p.stat().st_size > 0

            card_p = cards_dir / f"hinh_{f_slug}.md" if cards_dir.exists() else None
            card_ok = card_p and card_p.exists()

            if img_ok and card_ok:
                valid_figs += 1
            else:
                if not card_ok:
                    res.tickets.append(
                        DiagnosticTicket(
                            bundle_slug=bundle_dir.name,
                            error_category="MULTIMODAL_ASSET",
                            error_code="ERR_MISSING_FIGURE_CARD",
                            file_path=str((cards_dir / f"hinh_{f_slug}.md").relative_to(self.root_dir)) if cards_dir.exists() else "figures/cards/",
                            location=f"Figure '{tag}'",
                            detail=f"Figure card 'hinh_{f_slug}.md' is missing.",
                            assigned_skill="modernize_annex_engine",
                            action_required="Generate visual card for figure.",
                        )
                    )

        total = len(fig_list)
        rate = (valid_figs / total) * 100.0 if total > 0 else 100.0
        if stray_binaries:
            rate = min(rate, 50.0)
        res.p_multimodal = round(rate, 2)
        res.details["multimodal"] = {"total_figures": total, "valid_figures": valid_figs}

    def _check_structural_ast(self, bundle_dir: Path, res: BundleParityResult) -> None:
        """Check Structural AST & Atomic Form Templates."""
        clauses_file = bundle_dir / "clauses.json"
        has_clauses = clauses_file.exists()
        clauses_valid = False
        if has_clauses:
            try:
                c_data = json.loads(clauses_file.read_text(encoding="utf-8"))
                clauses_valid = isinstance(c_data, (dict, list)) and len(c_data) > 0
            except Exception:
                clauses_valid = False

        templates_dir = bundle_dir / "templates"
        templates_valid = True
        template_tickets = []
        if templates_dir.exists():
            t_files = list(templates_dir.rglob("*.md"))
            if not t_files:
                templates_valid = False
                template_tickets.append(("templates/", "Empty templates/ directory detected."))
            for tf in t_files:
                try:
                    t_text = tf.read_text(encoding="utf-8")
                    if not t_text.strip():
                        templates_valid = False
                        template_tickets.append((str(tf.relative_to(self.root_dir)), "Empty template file detected."))
                        continue

                    # Check for corrupted placeholder in frontmatter title or Markdown headings
                    lines = t_text.splitlines()
                    for l in lines[:25]:
                        l_strip = l.strip()
                        if l_strip.startswith("title:") or l_strip.startswith("#"):
                            if re.search(r"\.{4,}|……|___", l_strip):
                                templates_valid = False
                                template_tickets.append((str(tf.relative_to(self.root_dir)), f"Corrupted placeholder in title: '{l_strip[:40]}'"))
                                break

                    # Check for broken flattened table in template
                    if re.search(r"(?:__TT__|\bTT\b)\s*\n\s*\n\s*__(?:Danh mục|Tên sản phẩm)", t_text):
                        if "|" not in t_text:
                            templates_valid = False
                            template_tickets.append((str(tf.relative_to(self.root_dir)), "Flattened table detected in template. Must be 2D GFM Table."))
                except Exception:
                    pass

        score = 100.0
        if not clauses_valid:
            score -= 50.0
            res.tickets.append(
                DiagnosticTicket(
                    bundle_slug=bundle_dir.name,
                    error_category="STRUCTURAL_AST",
                    error_code="ERR_CLAUSES_AST_INVALID",
                    file_path=str(clauses_file.relative_to(self.root_dir)) if has_clauses else "clauses.json",
                    location="clauses.json",
                    detail="Missing or invalid clauses.json AST file.",
                    assigned_skill="ccba-markdown-document-processing",
                    action_required="Rebuild clauses.json AST hierarchy.",
                )
            )

        if not templates_valid:
            score -= 50.0
            for tf_path, msg in template_tickets[:2]:
                res.tickets.append(
                    DiagnosticTicket(
                        bundle_slug=bundle_dir.name,
                        error_category="STRUCTURAL_AST",
                        error_code="ERR_TEMPLATE_PLACEHOLDER",
                        file_path=tf_path,
                        location="templates/",
                        detail=msg,
                        assigned_skill="form-template-cleaner",
                        action_required="Clean placeholders and restore formal headings via form-template-cleaner.",
                    )
                )

        res.p_structure = max(0.0, round(score, 2))
        res.details["structure"] = {
            "clauses_valid": clauses_valid,
            "templates_valid": templates_valid,
        }


def format_markdown_report(results: List[BundleParityResult], duration_sec: float) -> str:
    """Render comprehensive Markdown report for Ground Truth Parity."""
    passed_count = sum(1 for r in results if r.passed)
    total_count = len(results)
    timestamp = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")

    lines = [
        "# Báo Cáo Đối Soát 1-1 Xác Định (Ground Truth Parity Report - Hardened v2.0)",
        "",
        f"- **Thời gian quét:** `{timestamp}`",
        f"- **Thời lượng:** `{duration_sec:.2f} giây`",
        f"- **Tổng số văn bản đã kiểm tra:** `{total_count}`",
        f"- **Số văn bản đạt chuẩn (Passed):** `{passed_count}/{total_count}` ({passed_count/total_count*100:.1f}%)",
        "",
        "---",
        "",
        "## 1. Bảng Tổng Hợp Vector Parity Score (VPS)",
        "",
        "| STT | Văn Bản (Slug) | Verbatim ($P_1$) | Table ($P_2$) | Math ($P_3$) | Multimodal ($P_4$) | Structure ($P_5$) | Trạng Thái |",
        "| :-: | :--- | :-: | :-: | :-: | :-: | :-: | :-: |",
    ]

    for idx, r in enumerate(results, 1):
        status = "✅ PASS" if r.passed else "❌ FAIL"
        lines.append(
            f"| {idx} | `{r.slug}` | {r.p_verbatim:.1f}% | {r.p_table:.1f}% | {r.p_math:.1f}% | {r.p_multimodal:.1f}% | {r.p_structure:.1f}% | {status} |"
        )

    all_tickets: List[DiagnosticTicket] = []
    for r in results:
        all_tickets.extend(r.tickets)

    lines.extend([
        "",
        "---",
        "",
        f"## 2. Danh Sách Thẻ Chẩn Đoán Lỗi (Diagnostic Error Tickets) - Tổng số: `{len(all_tickets)}`",
        "",
    ])

    if not all_tickets:
        lines.append("🎉 **Tuyệt vời! Không phát hiện bất kỳ lỗi hay độ lệch nào trên các văn bản đã kiểm tra.**")
    else:
        lines.append("| Văn Bản | Trục Lỗi | Mã Lỗi | Vị Trí | Chi Tiết | Skill Phụ Trách | Hành Động Yêu Cầu |")
        lines.append("| :--- | :--- | :--- | :--- | :--- | :--- | :--- |")
        for t in all_tickets:
            lines.append(
                f"| `{t.bundle_slug}` | **{t.error_category}** | `{t.error_code}` | {t.location} | {t.detail[:80]} | `{t.assigned_skill}` | {t.action_required[:80]} |"
            )

    lines.extend([
        "",
        "---",
        "",
        "## 3. Khuyến Nghị Cho Kỹ Sư & Nightly Tuner",
        "",
        "1. **Tự Động Phẫu Thuật (Surgical Fix):** Sử dụng các `DiagnosticTicket` ở trên để kích hoạt đúng skill chịu trách nhiệm.",
        "2. **Khóa An Toàn Hồi Quy (Regression Lock):** Chạy lại `python scripts/validate_legal_spoke.py` sau mỗi lần hiệu chỉnh để bảo đảm 100% không suy giảm chất lượng toàn kho.",
        "",
    ])

    return "\n".join(lines)


def main() -> int:
    parser = argparse.ArgumentParser(description="CCBA Legal Ground Truth Parity Verifier")
    parser.add_argument(
        "--cohorts",
        type=str,
        default="golden",
        help="Selection: 'golden' (5 Golden Cohorts), 'all' (all 55 bundles), or comma-separated slugs.",
    )
    parser.add_argument(
        "--json-out",
        type=str,
        default="",
        help="Optional path to output JSON report file.",
    )
    parser.add_argument(
        "--md-out",
        type=str,
        default="",
        help="Optional path to output Markdown summary report file.",
    )
    args = parser.parse_args()

    verifier = GroundTruthParityVerifier()

    slugs_to_test: List[str] = []
    if args.cohorts == "golden":
        for c in GOLDEN_COHORTS.values():
            slugs_to_test.extend(c["slugs"])
    elif args.cohorts == "all":
        for cat in ["01_vbpl", "02_qcvn", "03_tcvn"]:
            cat_dir = LEGAL_DOCS_DIR / cat
            if cat_dir.exists():
                for d in sorted(cat_dir.iterdir()):
                    if d.is_dir() and not d.name.startswith("."):
                        slugs_to_test.append(d.name)
    else:
        slugs_to_test = [s.strip() for s in args.cohorts.split(",") if s.strip()]

    print("=================================================================")
    print("  CCBA LEGAL DETERMINISTIC GROUND TRUTH PARITY ENGINE (v2.0)     ")
    print("=================================================================")
    print(f"Target Slugs Count : {len(slugs_to_test)}")
    print(f"Cohorts Mode       : {args.cohorts}")
    print("-----------------------------------------------------------------")

    start_time = time.time()
    results: List[BundleParityResult] = []

    for idx, slug in enumerate(slugs_to_test, 1):
        print(f"[{idx:02d}/{len(slugs_to_test):02d}] Auditing '{slug}'...", end=" ", flush=True)
        res = verifier.verify_bundle(slug)
        results.append(res)
        status = "PASSED" if res.passed else "FAILED"
        print(f"[{status}] VPS: {res.vps_tuple}")

    duration = time.time() - start_time
    print("-----------------------------------------------------------------")
    passed_total = sum(1 for r in results if r.passed)
    print(f"Summary: {passed_total}/{len(results)} Passed in {duration:.2f}s")
    print("=================================================================")

    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    t_str = datetime.now().strftime("%Y%m%d_%H%M%S")

    json_path = Path(args.json_out) if args.json_out else REPORTS_DIR / f"parity_report_{t_str}.json"
    md_path = Path(args.md_out) if args.md_out else REPORTS_DIR / f"parity_report_{t_str}.md"

    report_dict = {
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "duration_seconds": round(duration, 2),
        "total_documents": len(results),
        "passed_documents": passed_total,
        "results": [
            {
                "slug": r.slug,
                "category": r.category,
                "passed": r.passed,
                "vps": {
                    "p_verbatim": r.p_verbatim,
                    "p_table": r.p_table,
                    "p_math": r.p_math,
                    "p_multimodal": r.p_multimodal,
                    "p_structure": r.p_structure,
                },
                "details": r.details,
                "tickets": [asdict(t) for t in r.tickets],
            }
            for r in results
        ],
    }
    json_path.write_text(json.dumps(report_dict, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"JSON Report written to: {json_path}")

    md_content = format_markdown_report(results, duration)
    md_path.write_text(md_content, encoding="utf-8")
    print(f"Markdown Report written to: {md_path}")

    return 0 if passed_total == len(results) else 1


if __name__ == "__main__":
    sys.exit(main())
