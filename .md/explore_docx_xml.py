import zipfile
import xml.etree.ElementTree as ET
import docx

docx_path = r"d:\GitHubProjects\ccba-legal-knowledge\.md\extracted_docs\qcvn_06_2022_bxd\sua_doi_1_2023_qcvn_06_2022_bxd.docx"
md_path = r"d:\GitHubProjects\ccba-legal-knowledge\legal_docs\02_qcvn\qcvn_06_2022_bxd\sua_doi_1_2023_qcvn_06_2022_bxd.md"

print(f"=== DOCX ZIP CONTENTS ===")
with zipfile.ZipFile(docx_path, 'r') as z:
    for filename in z.namelist():
        info = z.getinfo(filename)
        print(f"{filename:35s} | size: {info.file_size:7d} bytes")

doc = docx.Document(docx_path)
print(f"\n=== PYTHON-DOCX PARSING ===")
print(f"Total doc.paragraphs: {len(doc.paragraphs)}")
non_empty_p = [p for p in doc.paragraphs if p.text.strip()]
print(f"Non-empty paragraphs: {len(non_empty_p)}")
print(f"Total doc.tables: {len(doc.tables)}")
for i, tbl in enumerate(doc.tables):
    rows = len(tbl.rows)
    cols = len(tbl.columns) if rows > 0 else 0
    print(f"  Table {i+1}: {rows} rows x {cols} columns")

# Check direct XML parsing of word/document.xml
with zipfile.ZipFile(docx_path, 'r') as z:
    doc_xml = z.read('word/document.xml')
    root = ET.fromstring(doc_xml)
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    all_p = root.findall('.//w:p', ns)
    all_tbl = root.findall('.//w:tbl', ns)
    print(f"\n=== DIRECT XML NODES IN document.xml ===")
    print(f"Total w:p elements: {len(all_p)}")
    print(f"Total w:tbl elements: {len(all_tbl)}")
    
    # Check if footnotes.xml exists
    if 'word/footnotes.xml' in z.namelist():
        fn_xml = z.read('word/footnotes.xml')
        fn_root = ET.fromstring(fn_xml)
        fn_p = fn_root.findall('.//w:p', ns)
        print(f"Footnotes xml w:p elements: {len(fn_p)}")
        for p in fn_p:
            texts = [t.text for t in p.findall('.//w:t', ns) if t.text]
            if texts:
                print(f"  FN P: {''.join(texts)}")
    else:
        print("No word/footnotes.xml")
