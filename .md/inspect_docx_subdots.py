"""
Check exact DOCX text for 2.1.1, 2.2.1, 5.1.1, 6.2.1
"""
import sys
import docx
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

DOCX_PATH = Path(r"d:\GitHubProjects\ccba-legal-knowledge\.md\extracted_docs\qcvn_06_2022_bxd\qcvn_06_2022_bxd.docx")
doc = docx.Document(str(DOCX_PATH))

indices = [338, 339, 340, 341, 342, 343, 390, 391, 392, 393, 394, 395, 396, 397, 898, 899, 900, 901, 902, 903, 904, 905, 1083, 1084, 1085, 1086, 1087, 1088, 1089, 1090, 1095]

for idx in indices:
    if idx < len(doc.paragraphs):
        p = doc.paragraphs[idx]
        print(f"Docx P#{idx:4d}: {repr(p.text[:120])}")
