"""Deep structural and content audit of sua_doi_1_2023_qcvn_06_2022_bxd.md."""

import json
import re
import sys
from pathlib import Path
from docx import Document

sys.stdout.reconfigure(encoding="utf-8")

bundle_dir = Path(__file__).resolve().parent.parent / "legal_docs" / "02_qcvn" / "qcvn_06_2022_bxd"
md_path = bundle_dir / "sua_doi_1_2023_qcvn_06_2022_bxd.md"
base_md_path = bundle_dir / "qcvn_06_2022_bxd.md"
docx_path = Path(__file__).resolve().parent.parent / ".md" / "extracted_docs" / "qcvn_06_2022_bxd" / "sua_doi_1_2023_qcvn_06_2022_bxd.docx"

md_text = md_path.read_text(encoding="utf-8")
md_lines = md_text.splitlines()

print("=================================================================")
print("     DEEP QUALITY & CONTENT AUDIT: SỬA ĐỔI 1:2023 QCVN 06       ")
print("=================================================================")

# 1. Check Document Header & Title
print("\n--- 1. TIÊU ĐỀ & MỞ ĐẦU ---")
for line in md_lines[:15]:
    if line.strip():
        print("  ", line.strip()[:100])

# 2. Check Headings Hierarchy
print("\n--- 2. CẤU TRÚC ĐỀ MỤC (HEADINGS HIERARCHY) ---")
h_counts = {"H1": 0, "H2": 0, "H3": 0, "H4": 0, "H5": 0, "H6": 0}
headings = []
for idx, line in enumerate(md_lines):
    m = re.match(r"^(#{1,6})\s+(.+)$", line.strip())
    if m:
        level = len(m.group(1))
        h_counts[f"H{level}"] += 1
        headings.append((level, m.group(2).strip(), idx + 1))

print(f"Tổng số tiêu đề: {len(headings)}")
print(f"Phân bổ cấp độ : {h_counts}")

# 3. Check Clauses Modified / Added
print("\n--- 3. CÁC ĐIỂM SỬA ĐỔI / BỔ SUNG (MODIFICATION POINTS) ---")
mod_points = []
for lvl, title, line_no in headings:
    if any(k in title for k in ["Sửa đổi", "Bổ sung", "Thay thế", "Bãi bỏ", "Hủy bỏ"]):
        mod_points.append((lvl, title, line_no))

print(f"Tổng số mục Sửa đổi / Bổ sung / Thay thế nhận diện được: {len(mod_points)}")
for lvl, title, line_no in mod_points[:10]:
    print(f"  Line {line_no:4d} [H{lvl}]: {title[:80]}")
if len(mod_points) > 10:
    print(f"  ... và {len(mod_points) - 10} mục khác.")

# 4. Check Tables in Sửa đổi 1
print("\n--- 4. BẢNG BIỂU KỸ THUẬT (TABLES IN SỬA ĐỔI 1) ---")
table_headings = [h for h in headings if "Bảng" in h[1] or "bảng" in h[1]]
print(f"Số đề mục Bảng trong Sửa đổi 1: {len(table_headings)}")
for lvl, title, line_no in table_headings:
    print(f"  Line {line_no:4d} [H{lvl}]: {title}")

table_pipe_blocks = 0
in_table = False
for line in md_lines:
    if line.strip().startswith("|"):
        if not in_table:
            table_pipe_blocks += 1
            in_table = True
    else:
        in_table = False
print(f"Số khối bảng GFM Pipe thực tế: {table_pipe_blocks}")

# 5. Check Cross-links to Base Document
print("\n--- 5. LIÊN KẾT ĐỐI SOÁT CHÉO (CROSS-LINKS TO BASE DOC) ---")
base_text = base_md_path.read_text(encoding="utf-8")
base_anchors = set(re.findall(r'<a id="([^"]+)"', base_text))

all_links = re.findall(r'\[([^\]]+)\]\(([^)]+)\)', md_text)
print(f"Tổng số liên kết trong file: {len(all_links)}")

valid_links = 0
broken_links = []
for label, target in all_links:
    if target.startswith("qcvn_06_2022_bxd.md#"):
        anchor = target.split("#", 1)[1]
        if anchor in base_anchors:
            valid_links += 1
        else:
            broken_links.append((label, target))
    else:
        print(f"  Khác: [{label}] -> {target}")

print(f"Liên kết trỏ về bản gốc hợp lệ: {valid_links} / {len(all_links)}")
if broken_links:
    print(f"❌ CÓ {len(broken_links)} LINK HỎNG:")
    for l, t in broken_links:
        print(f"    [{l}] -> {t}")
else:
    print("✅ 100% Liên kết trỏ về bản gốc đều chính xác nơi đích!")

# 6. Comparison with raw DOCX
print("\n--- 6. ĐỐI SOÁT VỚI FILE DOCX GỐC BAN HÀNH (THÔNG TƯ 09/2023/TT-BXD) ---")
if docx_path.exists():
    doc = Document(docx_path)
    docx_paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    print(f"Tổng số đoạn văn trong DOCX gốc: {len(docx_paras)}")
    print(f"Tổng số bảng trong DOCX gốc    : {len(doc.tables)}")

    # Strip Markdown syntax, HTML tags and URLs from MD for true text comparison
    clean_md_pure = re.sub(r"<[^>]+>", " ", md_text)
    clean_md_pure = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", clean_md_pure)
    clean_md_pure = re.sub(r"[_*#`><\"“”\'\(\)\[\]–—\-\.\,\:\;]", "", clean_md_pure)
    clean_md_pure = re.sub(r"\s+", " ", clean_md_pure).lower()

    matched_paras = 0
    missing_docx_paras = []

    for p in docx_paras:
        p_clean = re.sub(r"[_*#`><\"“”\'\(\)\[\]–—\-\.\,\:\;]", "", p).strip().lower()
        p_clean = re.sub(r"\s+", " ", p_clean)
        p_sample = p_clean[:35]
        if p_sample and p_sample in clean_md_pure:
            matched_paras += 1
        else:
            if len(p) > 15:
                missing_docx_paras.append(p)

    match_rate = (matched_paras / len(docx_paras) * 100) if docx_paras else 100
    print(f"Tỷ lệ khớp nội dung DOCX -> MD : {matched_paras}/{len(docx_paras)} ({match_rate:.2f}%)")
    if missing_docx_paras:
        print(f"Số đoạn chưa match chính xác ({len(missing_docx_paras)}):")
        for p in missing_docx_paras[:5]:
            print(f"  - {p[:90]}...")

print("=================================================================")
