"""Tests for scripts/docx_converter.py."""

from pathlib import Path
import pytest
import docx
from scripts.docx_converter import convert_docx_to_okf_bundle, normalize_docx_markdown

def create_dummy_docx(docx_path: Path) -> Path:
    """Helper to create a small valid .docx file."""
    doc = docx.Document()
    doc.add_heading("Điều 1. Phạm vi điều chỉnh", level=1)
    doc.add_paragraph("1. Quy định này điều chỉnh quản lý hoạt động xây dựng.")
    doc.add_heading("Điều 2. Đối tượng áp dụng", level=1)
    doc.add_paragraph("1. Áp dụng cho cơ quan, tổ chức, cá nhân có liên quan.")
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))
    return docx_path

def test_normalize_docx_markdown():
    raw_text = r"### ### __Điều 1. Phạm vi__ \(1.1\)"
    normalized = normalize_docx_markdown(raw_text)
    assert "### " in normalized
    assert r"\(" not in normalized

def test_convert_docx_to_okf_bundle_default_output(tmp_path: Path):
    input_docx = tmp_path / "test_doc.docx"
    create_dummy_docx(input_docx)

    bundle_dir = tmp_path / "bundle_test"
    res = convert_docx_to_okf_bundle(input_docx, bundle_dir)

    assert res.get("status") == "success"
    target_md = bundle_dir / "bundle_test.md"
    assert target_md.exists()
    assert (bundle_dir / "clauses.json").exists()
    assert (bundle_dir / "qa_benchmark.json").exists()

def test_convert_docx_to_okf_bundle_custom_output(tmp_path: Path):
    input_docx = tmp_path / "test_doc2.docx"
    create_dummy_docx(input_docx)

    bundle_dir = tmp_path / "bundle_test2"
    res = convert_docx_to_okf_bundle(input_docx, bundle_dir, output_filename="custom_name.md", doc_type="vbpl")

    assert res.get("status") == "success"
    target_md = bundle_dir / "custom_name.md"
    assert target_md.exists()
