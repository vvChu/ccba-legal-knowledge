"""Shared Pytest Fixtures & High-Performance Parser Models for E2E Tests.

Provides session-scoped parsed models of DOCX sources and Markdown OKF bundles,
enabling fast execution (< 2 seconds) of all 230+ assertions across 4 tiers.
"""

import csv
import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Set, Tuple

import docx
import pytest
import yaml

# =====================================================================
# 1. DATACLASS MODELS FOR HIGH-SPEED IN-MEMORY INDEXING
# =====================================================================

@dataclass
class DocxTableData:
    """Represents an extracted canonical technical table from DOCX."""
    table_idx: int
    num: str
    title: str
    slug: str
    headers: List[str]
    rows: List[List[str]]
    total_rows: int
    total_cols: int
    total_cells: int
    footnotes: List[str] = field(default_factory=list)

@dataclass
class DocxParsedBundle:
    """Represents a fully parsed official .docx document."""
    raw_paragraphs: List[str]
    body_paragraphs: List[str]
    headings: List[str]
    heading_map: Dict[str, str]
    tables: List[DocxTableData]
    table_map: Dict[str, DocxTableData]
    chapter_paragraphs: Dict[str, List[str]]

@dataclass
class MarkdownParsedBundle:
    """Represents a fully parsed OKF Markdown document."""
    raw_lines: List[str]
    raw_text: str
    normalized_text: str
    headings: List[Tuple[int, str, int]]  # (level, text, line_idx_1_based)
    heading_map: Dict[str, Tuple[int, str, int]]
    anchors: Set[str]
    pipe_tables: List[Dict[str, Any]]
    pipe_table_map: Dict[str, Dict[str, Any]]
    substantive_paragraphs: List[str]

# =====================================================================
# 2. HIGH-PRECISION PARSER IMPLEMENTATIONS
# =====================================================================

def parse_docx_qcvn(docx_path: Path) -> DocxParsedBundle:
    """Parses qcvn_06_2022_bxd.docx into structured paragraphs, headings, and 64 tables."""
    if not docx_path.exists():
        raise FileNotFoundError(f"Source DOCX not found at {docx_path}")

    doc = docx.Document(docx_path)
    raw_paras = [p.text.strip() for p in doc.paragraphs]
    non_empty_paras = [p for p in raw_paras if p]

    # Paragraphs 0..21 in DOCX contain title and TOC ('MỤC LỤC' down to 'PHỤ LỤC I').
    # Substantive body starts at paragraph 22 ('Lời nói đầu' / 'QUY CHUẨN KỸ THUẬT QUỐC GIA...').
    body_paras = [p.text.strip() for p in doc.paragraphs[22:] if p.text.strip()]

    # Heading patterns in body
    heading_pattern = re.compile(
        r"^(?:[0-9]+(?:\.[0-9]+)*|PHỤ LỤC\s+[A-Z]|CHƯƠNG\s+[0-9]+|Bảng\s+[A-Z0-9]+|\bMỤC\b)",
        re.IGNORECASE,
    )
    headings = []
    heading_map = {}
    for p in body_paras:
        if heading_pattern.match(p):
            headings.append(p)
            num_match = re.match(r"^([A-Z0-9]+(?:\.[0-9]+)*|PHỤ LỤC\s+[A-Z]|CHƯƠNG\s+[0-9]+)", p, re.IGNORECASE)
            if num_match:
                heading_map[num_match.group(1).lower()] = p

    # Chapter and Appendix grouping
    chapter_paragraphs: Dict[str, List[str]] = {
        "1": [], "2": [], "3": [], "4": [], "5": [], "6": [], "7": [],
        "A": [], "B": [], "C": [], "D": [], "E": [], "F": [], "G": [], "H": [], "I": [],
    }
    current_section = "1"
    for p in body_paras:
        m_chap = re.match(r"^([1-7])\s+[A-ZÀ-Ỹ\s]+$", p)
        m_app = re.match(r"^PHỤ LỤC\s+([A-I])\b", p, re.IGNORECASE)
        if m_chap:
            current_section = m_chap.group(1)
        elif m_app:
            current_section = m_app.group(1).upper()
        if current_section in chapter_paragraphs:
            chapter_paragraphs[current_section].append(p)

    # 64 Table Extraction & Titles
    title_pattern = re.compile(
        r"^(?:Bảng|Table)\s+([A-Z0-9]+(?:\.[0-9]+[a-z]?)?)\s*[-–:]\s*(.+)$",
        re.IGNORECASE,
    )
    table_titles = []
    for p_idx, text in enumerate(non_empty_paras):
        match = title_pattern.match(text)
        if match:
            table_titles.append({
                "num": match.group(1),
                "title": f"Bảng {match.group(1)} - {match.group(2).strip()}",
                "p_idx": p_idx,
            })

    tables: List[DocxTableData] = []
    table_map: Dict[str, DocxTableData] = {}

    for idx, table in enumerate(doc.tables):
        t_info = table_titles[idx] if idx < len(table_titles) else {
            "num": str(idx + 1),
            "title": f"Bảng {idx + 1}",
        }
        num_str = t_info["num"]
        slug = f"bang_{num_str.lower().replace('.', '_')}"
        if re.match(r"^\d+$", num_str):
            slug = f"bang_{int(num_str):02d}"

        grid: List[List[str]] = []
        footnotes: List[str] = []
        total_cells = 0

        for row in table.rows:
            row_cells = [c.text.replace("\n", " ").strip() for c in row.cells]
            total_cells += len(row_cells)

            # Detect footnote rows (single repeated text across all merged cells)
            if len(set(row_cells)) == 1 and (
                row_cells[0].startswith("CHÚ THÍCH")
                or re.match(r"^\d+\)\s+", row_cells[0])
                or len(row_cells[0]) > 80
            ):
                if row_cells[0] and row_cells[0] not in footnotes:
                    footnotes.append(row_cells[0])
                continue

            grid.append(row_cells)

        headers = grid[0] if grid else []
        rows = grid[1:] if len(grid) > 1 else []

        table_data = DocxTableData(
            table_idx=idx + 1,
            num=num_str,
            title=t_info["title"],
            slug=slug,
            headers=headers,
            rows=rows,
            total_rows=len(grid),
            total_cols=len(headers) if headers else 0,
            total_cells=total_cells,
            footnotes=footnotes,
        )
        tables.append(table_data)
        table_map[slug] = table_data
        table_map[num_str.lower()] = table_data

    return DocxParsedBundle(
        raw_paragraphs=raw_paras,
        body_paragraphs=body_paras,
        headings=headings,
        heading_map=heading_map,
        tables=tables,
        table_map=table_map,
        chapter_paragraphs=chapter_paragraphs,
    )

def parse_markdown_qcvn(md_path: Path) -> MarkdownParsedBundle:
    """Parses qcvn_06_2022_bxd.md into structured headings, anchors, and GFM pipe tables."""
    if not md_path.exists():
        raise FileNotFoundError(f"Target Markdown not found at {md_path}")

    raw_text = md_path.read_text(encoding="utf-8")
    bundle_dir = md_path.parent
    annexes_dir = bundle_dir / "annexes"
    if annexes_dir.exists():
        for annex_file in sorted(annexes_dir.glob("*.md")):
            raw_text += "\n\n" + annex_file.read_text(encoding="utf-8")

    raw_lines = raw_text.splitlines()
    normalized_text = re.sub(r"\s+", " ", raw_text).lower()

    headings: List[Tuple[int, str, int]] = []
    heading_map: Dict[str, Tuple[int, str, int]] = {}
    anchors: Set[str] = set()
    pipe_tables: List[Dict[str, Any]] = []
    pipe_table_map: Dict[str, Dict[str, Any]] = {}
    substantive_paragraphs: List[str] = []

    # Extract anchors (supports both standalone and inlined <a id="..." name="..."></a>)
    for anchor_m in re.finditer(r'<a\s+id="([^"]+)"', raw_text):
        anchors.add(anchor_m.group(1))

    # Parse line by line
    in_table = False
    current_table_lines: List[str] = []
    current_table_anchor = ""
    current_table_title = ""

    for line_idx, line in enumerate(raw_lines):
        stripped = line.strip()

        # Headings
        h_match = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if h_match:
            level = len(h_match.group(1))
            h_text_raw = h_match.group(2).strip()
            # Extract anchor if embedded in heading
            a_in_h = re.search(r'<a\s+id="([^"]+)"', h_text_raw)
            if a_in_h:
                current_table_anchor = a_in_h.group(1)
            # Clean heading text
            h_text = re.sub(r'<a\s+[^>]+></a>\s*', '', h_text_raw).strip()
            headings.append((level, h_text, line_idx + 1))
            heading_map[h_text.lower()] = (level, h_text, line_idx + 1)
            # Check if this heading is a table title
            if re.match(r"^Bảng\s+([A-Z0-9]+(?:\.[0-9]+[a-z]?)?)", h_text, re.IGNORECASE):
                current_table_title = h_text

        # Table detection
        if stripped.startswith("|") and stripped.endswith("|"):
            if not in_table:
                in_table = True
                current_table_lines = [stripped]
            else:
                current_table_lines.append(stripped)
        else:
            if in_table:
                in_table = False
                # Parse the completed pipe table
                if len(current_table_lines) >= 2:
                    raw_headers = [c.strip() for c in current_table_lines[0].strip("|").split("|")]
                    data_rows = []
                    for row_line in current_table_lines[2:]:  # skip separator line
                        row_cells = [c.strip() for c in row_line.strip("|").split("|")]
                        data_rows.append(row_cells)

                    t_dict = {
                        "title": current_table_title,
                        "anchor": current_table_anchor,
                        "headers": raw_headers,
                        "rows": data_rows,
                        "total_rows": len(current_table_lines),
                        "total_cols": len(raw_headers),
                        "line_start": line_idx - len(current_table_lines) + 1,
                    }
                    pipe_tables.append(t_dict)
                    if current_table_title:
                        pipe_table_map[current_table_title.lower()] = t_dict

                current_table_lines = []
                current_table_title = ""

            # Check if anchor precedes table
            a_m = re.match(r'<a\s+id="([^"]+)"></a>', stripped)
            if a_m:
                current_table_anchor = a_m.group(1)

            # Substantive text paragraphs
            if (
                stripped
                and not stripped.startswith("#")
                and not stripped.startswith("<a")
                and not stripped.startswith("---")
                and not stripped.startswith("```")
                and not stripped.startswith(">")
                and len(stripped) > 20
            ):
                substantive_paragraphs.append(stripped)

    return MarkdownParsedBundle(
        raw_lines=raw_lines,
        raw_text=raw_text,
        normalized_text=normalized_text,
        headings=headings,
        heading_map=heading_map,
        anchors=anchors,
        pipe_tables=pipe_tables,
        pipe_table_map=pipe_table_map,
        substantive_paragraphs=substantive_paragraphs,
    )

def parse_docx_amendment(docx_path: Path) -> Dict[str, Any]:
    """Parses sua_doi_1_2023_qcvn_06_2022_bxd.docx into paragraphs, directives, and tables."""
    if not docx_path.exists():
        raise FileNotFoundError(f"Source Amendment DOCX not found at {docx_path}")

    doc = docx.Document(docx_path)
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    # Directives identification
    directive_pattern = re.compile(
        r"^(?:\d+[\.\)]\s*)?(?:Sửa đổi|Bổ sung|Bãi bỏ|Thay thế)\b",
        re.IGNORECASE,
    )
    directives = [p for p in paras if directive_pattern.search(p)]

    # Tables
    tables_data = []
    for t_idx, table in enumerate(doc.tables):
        grid = []
        for row in table.rows:
            grid.append([c.text.replace("\n", " ").strip() for c in row.cells])
        tables_data.append({
            "table_idx": t_idx + 1,
            "grid": grid,
            "rows_count": len(grid),
            "cols_count": len(grid[0]) if grid else 0,
        })

    return {
        "paragraphs": paras,
        "total_paragraphs": len(paras),
        "directives": directives,
        "total_directives": len(directives),
        "tables": tables_data,
        "total_tables": len(tables_data),
    }

def parse_markdown_amendment(md_path: Path) -> Dict[str, Any]:
    """Parses sua_doi_1_2023_qcvn_06_2022_bxd.md into directives, anchors, links, and Table 10."""
    if not md_path.exists():
        raise FileNotFoundError(f"Target Amendment Markdown not found at {md_path}")

    raw_text = md_path.read_text(encoding="utf-8")
    raw_lines = raw_text.splitlines()

    # Directives
    directive_headers = [line.strip() for line in raw_lines if line.strip().startswith("####")]

    # Anchors
    anchors = set(re.findall(r'<a\s+id="([^"]+)"></a>', raw_text))

    # Cross-links to base QCVN
    cross_links = re.findall(r'\[([^\]]+)\]\(([^)]*qcvn_06_2022_bxd\.md[^)]*)\)', raw_text)

    # Table 10 extraction
    table_10_lines = []
    in_table_10 = False
    for line in raw_lines:
        s = line.strip()
        if "sd1-bang-10" in s or "Bảng 10" in s:
            in_table_10 = True
        if in_table_10:
            if s.startswith("|") and s.endswith("|"):
                table_10_lines.append(s)
            elif table_10_lines and not s.startswith("|"):
                break

    return {
        "raw_text": raw_text,
        "raw_lines": raw_lines,
        "directives": directive_headers,
        "total_directives": len(directive_headers),
        "anchors": anchors,
        "cross_links": cross_links,
        "table_10_lines": table_10_lines,
    }

# =====================================================================
# 3. SESSION-SCOPED PYTEST FIXTURES
# =====================================================================

@pytest.fixture(scope="session")
def repo_root() -> Path:
    """Project repository root directory."""
    return Path(__file__).resolve().parent.parent

@pytest.fixture(scope="session")
def qcvn_bundle_dir(repo_root: Path) -> Path:
    """QCVN 06:2022/BXD OKF Bundle directory."""
    return repo_root / "legal_docs" / "02_qcvn" / "qcvn_06_2022_bxd"

@pytest.fixture(scope="session")
def extracted_docs_dir(repo_root: Path) -> Path:
    """Extracted raw official .docx directory."""
    return repo_root / ".md" / "extracted_docs" / "qcvn_06_2022_bxd"

@pytest.fixture(scope="session")
def docx_qcvn_path(extracted_docs_dir: Path) -> Path:
    """Path to qcvn_06_2022_bxd.docx."""
    return extracted_docs_dir / "qcvn_06_2022_bxd.docx"

@pytest.fixture(scope="session")
def docx_sd1_path(extracted_docs_dir: Path) -> Path:
    """Path to sua_doi_1_2023_qcvn_06_2022_bxd.docx."""
    return extracted_docs_dir / "sua_doi_1_2023_qcvn_06_2022_bxd.docx"

@pytest.fixture(scope="session")
def md_qcvn_path(qcvn_bundle_dir: Path) -> Path:
    """Path to qcvn_06_2022_bxd.md."""
    return qcvn_bundle_dir / "qcvn_06_2022_bxd.md"

@pytest.fixture(scope="session")
def md_sd1_path(qcvn_bundle_dir: Path) -> Path:
    """Path to sua_doi_1_2023_qcvn_06_2022_bxd.md."""
    target = qcvn_bundle_dir / "sources" / "sua_doi_1_2023_qcvn_06_2022_bxd.md"
    if not target.exists():
        target = qcvn_bundle_dir / "sua_doi_1_2023_qcvn_06_2022_bxd.md"
    return target

@pytest.fixture(scope="session")
def qcvn_docx_parsed(docx_qcvn_path: Path) -> DocxParsedBundle:
    """Cached in-memory parsed model of qcvn_06_2022_bxd.docx."""
    return parse_docx_qcvn(docx_qcvn_path)

@pytest.fixture(scope="session")
def sd1_docx_parsed(docx_sd1_path: Path) -> Dict[str, Any]:
    """Cached in-memory parsed model of sua_doi_1_2023_qcvn_06_2022_bxd.docx."""
    return parse_docx_amendment(docx_sd1_path)

@pytest.fixture(scope="session")
def qcvn_md_parsed(md_qcvn_path: Path) -> MarkdownParsedBundle:
    """Cached in-memory parsed model of qcvn_06_2022_bxd.md."""
    return parse_markdown_qcvn(md_qcvn_path)

@pytest.fixture(scope="session")
def sd1_md_parsed(md_sd1_path: Path) -> Dict[str, Any]:
    """Cached in-memory parsed model of sua_doi_1_2023_qcvn_06_2022_bxd.md."""
    return parse_markdown_amendment(md_sd1_path)


@pytest.fixture(scope="session")
def json_tables_map(qcvn_bundle_dir: Path) -> Dict[str, dict]:
    """Map of slug/stem to JSON table dict for all 64 files in tables/json/."""
    tables_dir = qcvn_bundle_dir / "tables" / "json"
    result: Dict[str, dict] = {}
    if tables_dir.exists():
        for json_file in tables_dir.glob("*.json"):
            with open(json_file, "r", encoding="utf-8") as f:
                try:
                    data = json.load(f)
                    result[json_file.stem] = data
                    if "table_id" in data:
                        result[data["table_id"]] = data
                except (json.JSONDecodeError, OSError):
                    pass
    return result

@pytest.fixture(scope="session")
def csv_tables_map(qcvn_bundle_dir: Path) -> Dict[str, List[List[str]]]:
    """Map of slug/stem to list of CSV rows for all 64 files in tables/csv/."""
    tables_dir = qcvn_bundle_dir / "tables" / "csv"
    result: Dict[str, List[List[str]]] = {}
    if tables_dir.exists():
        for csv_file in tables_dir.glob("*.csv"):
            with open(csv_file, "r", encoding="utf-8") as f:
                try:
                    reader = csv.reader(f)
                    result[csv_file.stem] = list(reader)
                except (csv.Error, OSError):
                    pass
    return result

@pytest.fixture(scope="session")
def clauses_ast_data(qcvn_bundle_dir: Path) -> List[dict]:
    """Parsed content of clauses.json."""
    clauses_file = qcvn_bundle_dir / "clauses.json"
    if clauses_file.exists():
        with open(clauses_file, "r", encoding="utf-8") as f:
            try:
                return json.load(f)
            except (json.JSONDecodeError, OSError):
                return []
    return []

@pytest.fixture(scope="session")
def qa_benchmark_data(qcvn_bundle_dir: Path) -> List[dict]:
    """Parsed content of qa_benchmark.json."""
    qa_file = qcvn_bundle_dir / "qa_benchmark.json"
    if qa_file.exists():
        with open(qa_file, "r", encoding="utf-8") as f:
            try:
                return json.load(f)
            except (json.JSONDecodeError, OSError):
                return []
    return []

@pytest.fixture(scope="session")
def legal_registry_data(repo_root: Path) -> dict:
    """Parsed content of legal_registry.yaml."""
    reg_file = repo_root / "legal_registry.yaml"
    if reg_file.exists():
        with open(reg_file, "r", encoding="utf-8") as f:
            try:
                return yaml.safe_load(f) or {}
            except (yaml.YAMLError, OSError):
                return {}
    return {}

@pytest.fixture(scope="session")
def all_registered_bundles(repo_root: Path, legal_registry_data: dict) -> List[Dict[str, Any]]:
    """Returns list of all active registered bundle entries across categories."""
    bundles = []
    for doc in legal_registry_data.get("documents", []):
        bp = doc.get("bundle_path")
        if bp:
            bundle_dir = repo_root / bp.strip("/")
            if bundle_dir.exists():
                bundles.append({**doc, "absolute_dir": bundle_dir})
    return bundles

@pytest.fixture(scope="session")
def all_bundle_dirs(all_registered_bundles: List[Dict[str, Any]]) -> List[Path]:
    """Returns list of Path objects for all registered OKF bundle directories."""
    return [b["absolute_dir"] for b in all_registered_bundles]

